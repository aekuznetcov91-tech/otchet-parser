import pickle
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('scratch/cross_brand_data.pkl', 'rb') as f:
    data = pickle.load(f)

deals_by_client = data['deals_by_client']
partner_leads_by_client = data['partner_leads_by_client']
crm_leads_by_client = data['crm_leads_by_client']

aux_kw = ('КРЕДИТ', 'КАСКО', 'ОСАГО', 'ГАП', 'СТРАХОВ', 'СЕРТИФИКАТ', 'ВНЕСЕНИЕ', 'АВАНС', 'ОФОРМЛЕНИЕ', 'ДОГОВОР', 'УСЛУГА', 'КОМИССИЯ', 'ДОП')

def is_js(s):
    if not s: return False
    u = str(s).upper()
    return any(k in u for k in ('JETOUR', 'JETOR', 'ДЖЕТУР', 'SOUEAST', 'SOUEAS', 'СОУИСТ', 'DASHING', 'T2', 'T1', 'S07', 'S09', 'S06'))

wb = openpyxl.Workbook()
wb.remove(wb.active)

# Styling definitions
header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
header_fill_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid") # Dark amber
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
bold_font = Font(name="Calibri", size=10, bold=True)
regular_font = Font(name="Calibri", size=10)
link_font = Font(name="Calibri", size=10, color="0000EE", underline="single")
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# -------------------------------------------------------------
# TAB 1: CASE 2 (Утечка переданных лидов в другие бренды)
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Кейс 2 - Утечка лидов JS")
headers_case2 = [
    "Client ID", "CRM Ссылка", "Переданный бренд (Лид)", "Модель в лиде", 
    "Дилер (Партнер)", "Дата передачи", "Купленный бренд", "Купленный автомобиль", 
    "Дата сделки", "Сумма сделки (₽)", "Канал (B2C)", "Менеджер", "Номер сделки"
]
ws2.append(headers_case2)
for col_idx in range(1, len(headers_case2) + 1):
    cell = ws2.cell(row=1, column=col_idx)
    cell.fill = header_fill_amber
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws2.row_dimensions[1].height = 28

case2_rows = []
for cid, dlist in deals_by_client.items():
    car_deals = [d for d in dlist if not is_js(d.get('tovar')) and not is_js(d.get('brand')) and not any(kw in str(d.get('tovar')).upper() for kw in aux_kw)]
    if not car_deals: continue
    
    c_part = partner_leads_by_client.get(cid, [])
    c_crm = crm_leads_by_client.get(cid, [])
    
    js_transfers = []
    for l in c_part:
        if is_js(l.get('brand_raw')) or is_js(l.get('brand_norm')) or is_js(l.get('model')):
            js_transfers.append(l)
    for l in c_crm:
        if l.get('is_transferred') and (is_js(l.get('brand_raw')) or is_js(l.get('brand_norm')) or is_js(l.get('model'))):
            js_transfers.append(l)
            
    if js_transfers:
        for t in js_transfers:
            for deal in car_deals:
                case2_rows.append({
                    'cid': cid,
                    'lead_b': t.get('brand_norm') or t.get('brand_raw'),
                    'lead_m': t.get('model'),
                    'partner': t.get('partner') or 'Не указан',
                    'lead_dt': t.get('date_str'),
                    'deal_b': deal.get('brand'),
                    'deal_m': deal.get('tovar'),
                    'deal_dt': deal.get('date_str'),
                    'price': deal.get('price'),
                    'b2c': deal.get('b2c'),
                    'manager': deal.get('manager'),
                    'deal_id': deal.get('deal_id')
                })

df2_clean = pd.DataFrame(case2_rows).drop_duplicates(subset=['cid', 'deal_m', 'deal_id'])
df2_clean.sort_values(by=['lead_dt', 'cid'], ascending=False, inplace=True)

for r_idx, r in enumerate(df2_clean.itertuples(), start=2):
    crm_url = f"https://backoffice.x.sberauto.com/crm/manager/{r.cid}"
    row_data = [
        str(r.cid),
        crm_url,
        str(r.lead_b),
        str(r.lead_m or ''),
        str(r.partner),
        str(r.lead_dt),
        str(r.deal_b),
        str(r.deal_m),
        str(r.deal_dt),
        r.price,
        str(r.b2c),
        str(r.manager),
        str(r.deal_id)
    ]
    ws2.append(row_data)
    ws2.cell(row=r_idx, column=10).number_format = '#,##0 ₽'
    ws2.cell(row=r_idx, column=2).hyperlink = crm_url
    ws2.cell(row=r_idx, column=2).font = link_font
    for c_idx in range(1, len(row_data) + 1):
        c = ws2.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if c_idx != 2: c.font = regular_font
        if c_idx in (1, 6, 9, 11, 13): c.alignment = Alignment(horizontal="center")

# -------------------------------------------------------------
# TAB 2: CASE 1 (Кросс-продажи в JETOUR / SOUEAST)
# -------------------------------------------------------------
ws1 = wb.create_sheet(title="Кейс 1 - Покупки JETOUR&SOUEAST")
headers_case1 = [
    "Тип кросс-продажи", "Client ID", "CRM Ссылка", "Первоначальный бренд", 
    "Дата первого лида", "Все запрошенные бренды", "Купленный бренд", 
    "Купленный автомобиль", "Дата сделки", "Сумма сделки (₽)", "Канал (B2C)", "Менеджер", "Номер сделки"
]
ws1.append(headers_case1)
for col_idx in range(1, len(headers_case1) + 1):
    cell = ws1.cell(row=1, column=col_idx)
    cell.fill = header_fill_blue
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws1.row_dimensions[1].height = 28

case1_rows = []
for cid, dlist in deals_by_client.items():
    js_deals = [d for d in dlist if is_js(d.get('brand')) or is_js(d.get('tovar'))]
    if not js_deals: continue
    
    c_crm = crm_leads_by_client.get(cid, [])
    c_part = partner_leads_by_client.get(cid, [])
    all_leads = []
    for l in c_crm:
        b = l.get('brand_norm') or l.get('brand_raw')
        if b: all_leads.append(l)
    for l in c_part:
        b = l.get('brand_norm') or l.get('brand_raw')
        if b: all_leads.append(l)
        
    other_leads = [l for l in all_leads if not is_js(l.get('brand_norm') or l.get('brand_raw'))]
    js_leads = [l for l in all_leads if is_js(l.get('brand_norm') or l.get('brand_raw'))]
    
    if other_leads:
        valid_dated = [l for l in all_leads if l.get('date')]
        valid_dated.sort(key=lambda x: x['date'])
        first_lead = valid_dated[0] if valid_dated else None
        
        is_pure = (len(js_leads) == 0)
        cross_type = "Чистый Cross-Sell (0 лидов на JS)" if is_pure else "Кроссовер (смена бренда)"
        
        other_b_names = sorted(set((l.get('brand_norm') or l.get('brand_raw') or '').strip() for l in other_leads if l.get('brand_norm') or l.get('brand_raw')))
        
        for deal in js_deals:
            case1_rows.append({
                'cross_type': cross_type,
                'cid': cid,
                'first_b': (first_lead.get('brand_norm') or first_lead.get('brand_raw')) if first_lead else (other_b_names[0] if other_b_names else 'Другой'),
                'first_dt': first_lead.get('date_str') if first_lead else '',
                'all_other_b': ", ".join(other_b_names),
                'deal_b': deal.get('brand'),
                'deal_m': deal.get('tovar'),
                'deal_dt': deal.get('date_str'),
                'price': deal.get('price'),
                'b2c': deal.get('b2c'),
                'manager': deal.get('manager'),
                'deal_id': deal.get('deal_id')
            })

df1_clean = pd.DataFrame(case1_rows).drop_duplicates(subset=['cid', 'deal_m', 'deal_id'])
df1_clean.sort_values(by=['cross_type', 'deal_dt'], ascending=[True, False], inplace=True)

for r_idx, r in enumerate(df1_clean.itertuples(), start=2):
    crm_url = f"https://backoffice.x.sberauto.com/crm/manager/{r.cid}"
    row_data = [
        str(r.cross_type),
        str(r.cid),
        crm_url,
        str(r.first_b),
        str(r.first_dt),
        str(r.all_other_b),
        str(r.deal_b),
        str(r.deal_m),
        str(r.deal_dt),
        r.price,
        str(r.b2c),
        str(r.manager),
        str(r.deal_id)
    ]
    ws1.append(row_data)
    ws1.cell(row=r_idx, column=10).number_format = '#,##0 ₽'
    ws1.cell(row=r_idx, column=3).hyperlink = crm_url
    ws1.cell(row=r_idx, column=3).font = link_font
    for c_idx in range(1, len(row_data) + 1):
        c = ws1.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if c_idx != 3: c.font = regular_font
        if c_idx in (1, 2, 5, 9, 11, 13): c.alignment = Alignment(horizontal="center")

# Auto-fit columns for both sheets
for ws in [ws1, ws2]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

out_dir = os.path.expanduser('~/Downloads')
out_path = os.path.join(out_dir, 'Кросс_брендовый_анализ_Jetour_Soueast_Июнь_Сентябрь_2026.xlsx')
wb.save(out_path)
print(f"Successfully generated Excel report: {out_path}")
