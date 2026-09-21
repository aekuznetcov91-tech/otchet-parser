import os
import sys
import json
import datetime
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJECT_ROOT, 'scripts'))
from parser_engine import read_tabular_file, get_exact_val, parse_custom_date

OUTPUT_PATH = os.path.join(PROJECT_ROOT, 'raw_data', 'Продажи_GAC_Август_2026_Москва_и_РФ.xlsx')
DESKTOP_PATH = os.path.expanduser('~/Desktop/Продажи_GAC_Август_2026_Москва_и_РФ.xlsx')

def normalize_gac_model(m):
    m_up = (m or '').upper()
    if 'GS8' in m_up or 'CS8' in m_up or m_up == 'GL' or 'GS 8' in m_up:
        return 'GAC GS8'
    elif 'GS4' in m_up or 'CS4' in m_up or 'GS 4' in m_up:
        return 'GAC GS4'
    elif 'S7' in m_up:
        return 'GAC GS3 (S7)'
    elif 'S9' in m_up:
        return 'GAC M8 (S9)'
    elif 'HYPTEC' in m_up:
        return 'GAC Hyptec HT'
    elif 'AION' in m_up:
        return 'GAC Aion V'
    return m or 'Не указана'

def main():
    print("[*] Загрузка данных data.json...")
    data_json_path = os.path.join(PROJECT_ROOT, 'site', 'data.json')
    with open(data_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    pdb = data.get('sys_db_partners', [])
    gac_aug_pdb = [r for r in pdb if r.get('Type') == 'Сделка' and r.get('Brand') == 'GAC' and r.get('Month') == '2026-08']
    
    # Read raw deals for enriched fields (City, Company, etc.)
    deals_file = os.path.join(PROJECT_ROOT, 'raw_data', 'DEAL_20260904_c9102cd3_6a9a5fa93d07b.xls')
    if not os.path.exists(deals_file):
        deals_file = os.path.join(PROJECT_ROOT, 'raw_data', 'DEAL_20260918_d9bd51b9_6aacd9139ab68.xls')
    raw_deals = {}
    if os.path.exists(deals_file):
        for sname, rows in read_tabular_file(deals_file):
            for r in rows:
                did = str(get_exact_val(r, 'ID', 'IDСДЕЛКИ') or '').strip()
                if did and did not in raw_deals:
                    raw_deals[did] = r

    # Build deal list
    deals = []
    seen_dids = set()
    for r in gac_aug_pdb:
        did = str(r.get('DealId'))
        if did in seen_dids:
            continue
        seen_dids.add(did)
        raw_r = raw_deals.get(did, {})
        city = str(get_exact_val(raw_r, 'ГОРОДB2C', 'ГОРОД. B2C', 'ГОРОД') or '').strip() or 'Не указан'
        company = str(get_exact_val(raw_r, 'КОМПАНИЯ', 'КОМПАНИЯ: НАЗВАНИЕ КОМПАНИИ') or r.get('RawPartner') or '').strip()
        
        # Excel date to string
        d_val = r.get('Date')
        d_str = ""
        if d_val:
            try:
                dt = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(float(d_val)))
                d_str = dt.strftime('%d.%m.%Y')
            except Exception:
                d_str = str(d_val)
                
        model_raw = r.get('Model') or ''
        model_norm = normalize_gac_model(model_raw)
        
        deals.append({
            'deal_id': did,
            'date_str': d_str,
            'date_val': d_val,
            'partner': r.get('Partner') or 'Не указан',
            'company': company or r.get('RawPartner') or 'Не указан',
            'city': city,
            'brand': 'GAC',
            'model_raw': model_raw,
            'model_norm': model_norm,
            'vin': r.get('VIN') or '',
            'b2c': r.get('B2C') or 'Не указан',
            'price': float(r.get('Price') or 0),
            'comm': float(r.get('Comm') or 0),
            'manager': r.get('Manager') or 'Не указан',
            'kam': r.get('KAM') or 'Не указан',
            'client_id': r.get('ClientId') or '',
            'lead_id': r.get('LeadId') or ''
        })

    print(f"[*] Собрано сделок GAC за август: {len(deals)}")

    # Classify by Moscow / MO
    moscow_strictly = [d for d in deals if d['city'].lower() == 'москва']
    moscow_mo = [d for d in deals if d['city'].lower() == 'москва' or any(mo in d['city'].lower() for mo in ['одинцово', 'нагорное', 'балашиха', 'химки'])]

    print(f"[*] В Москве (строго): {len(moscow_strictly)}")
    print(f"[*] В Москве и МО: {len(moscow_mo)}")

    # Create Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # Colors & Styles
    c_dark_navy = "0F172A"
    c_sber_green = "21A038"
    c_slate_header = "1E293B"
    c_light_gray = "F8FAFC"
    c_zebra = "F1F5F9"
    c_border = "CBD5E1"
    c_accent_blue = "0284C7"

    f_title = Font(name="Calibri", size=16, bold=True, color="0F172A")
    f_subtitle = Font(name="Calibri", size=11, italic=True, color="64748B")
    f_kpi_num = Font(name="Calibri", size=18, bold=True, color="21A038")
    f_kpi_label = Font(name="Calibri", size=9, bold=True, color="64748B")
    f_head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    f_subhead = Font(name="Calibri", size=11, bold=True, color="0F172A")
    f_total = Font(name="Calibri", size=10, bold=True, color="0F172A")
    f_regular = Font(name="Calibri", size=10)
    f_bold = Font(name="Calibri", size=10, bold=True)

    fill_head = PatternFill(start_color=c_slate_header, end_color=c_slate_header, fill_type="solid")
    fill_subhead = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_kpi = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_zebra = PatternFill(start_color=c_zebra, end_color=c_zebra, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=c_border),
        right=Side(style='thin', color=c_border),
        top=Side(style='thin', color=c_border),
        bottom=Side(style='thin', color=c_border)
    )
    total_border = Border(
        left=Side(style='thin', color=c_border),
        right=Side(style='thin', color=c_border),
        top=Side(style='thin', color="0F172A"),
        bottom=Side(style='double', color="0F172A")
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # ==========================================
    # SHEET 1: Сводка_РФ_и_Москва
    # ==========================================
    ws1 = wb.create_sheet(title="Сводка_РФ_и_Москва")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G1")
    ws1["A1"] = "ОТЧЕТ ПО ПРОДАЖАМ АВТОМОБИЛЕЙ GAC (СБЕРАВТО)"
    ws1["A1"].font = f_title
    ws1["A1"].alignment = align_left

    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"Период: Август 2026 года | Данные обновлены: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws1["A2"].font = f_subtitle
    ws1["A2"].alignment = align_left

    # KPI Cards
    kpis = [
        ("ВСЕГО ПРОДАЖ GAC (РФ)", len(deals), "#,##0 шт."),
        ("ИЗ НИХ ОПТ (МП2)", sum(1 for d in deals if 'МП2' in d['b2c'].upper()), "#,##0 шт."),
        ("ИЗ НИХ РОЗНИЦА", sum(1 for d in deals if 'МП2' not in d['b2c'].upper()), "#,##0 шт."),
        ("ПРОДАЖИ В МОСКВЕ", len(moscow_strictly), "#,##0 шт."),
        ("МОСКВА + МО", len(moscow_mo), "#,##0 шт."),
        ("СУММА КОМИССИИ РФ", sum(d['comm'] for d in deals), "#,##0.00 ₽"),
        ("ОБЩИЙ ОБОРОТ (РФ)", sum(d['price'] for d in deals), "#,##0.00 ₽"),
    ]

    for col_idx, (kpi_label, kpi_val, num_fmt) in enumerate(kpis, start=1):
        c_letter = get_column_letter(col_idx)
        ws1[f"{c_letter}4"] = kpi_label
        ws1[f"{c_letter}4"].font = f_kpi_label
        ws1[f"{c_letter}4"].alignment = align_center
        ws1[f"{c_letter}4"].fill = fill_kpi
        ws1[f"{c_letter}4"].border = thin_border

        ws1[f"{c_letter}5"] = kpi_val
        ws1[f"{c_letter}5"].font = f_kpi_num
        ws1[f"{c_letter}5"].alignment = align_center
        ws1[f"{c_letter}5"].fill = fill_kpi
        ws1[f"{c_letter}5"].border = thin_border
        ws1[f"{c_letter}5"].number_format = num_fmt

    # Table 1: Model breakdown across Russia
    ws1["A7"] = "1. Продажи GAC по всей России в разрезе моделей"
    ws1["A7"].font = f_subhead
    
    headers_t1 = ["Модель", "Сделок (шт.)", "Доля рынка GAC", "Сумма продаж (₽)", "Комиссия СберАвто (₽)", "Ср. чек (₽)", "Ср. комиссия (₽)"]
    for c_i, h in enumerate(headers_t1, 1):
        cell = ws1.cell(row=8, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    model_agg = defaultdict(lambda: {'count': 0, 'price': 0.0, 'comm': 0.0})
    for d in deals:
        m = d['model_norm']
        model_agg[m]['count'] += 1
        model_agg[m]['price'] += d['price']
        model_agg[m]['comm'] += d['comm']

    sorted_models = sorted(model_agg.items(), key=lambda x: x[1]['count'], reverse=True)
    r_idx = 9
    tot_cnt = len(deals)
    for m, inf in sorted_models:
        ws1.cell(row=r_idx, column=1, value=m).alignment = align_left
        ws1.cell(row=r_idx, column=2, value=inf['count']).number_format = "#,##0"
        ws1.cell(row=r_idx, column=3, value=inf['count'] / tot_cnt).number_format = "0.0%"
        ws1.cell(row=r_idx, column=4, value=inf['price']).number_format = "#,##0.00 ₽"
        ws1.cell(row=r_idx, column=5, value=inf['comm']).number_format = "#,##0.00 ₽"
        ws1.cell(row=r_idx, column=6, value=inf['price'] / inf['count'] if inf['count'] else 0).number_format = "#,##0.00 ₽"
        ws1.cell(row=r_idx, column=7, value=inf['comm'] / inf['count'] if inf['count'] else 0).number_format = "#,##0.00 ₽"
        
        for col in range(1, 8):
            c = ws1.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if col in [2, 3]: c.alignment = align_center
            elif col >= 4: c.alignment = align_right
        r_idx += 1

    # Total row Table 1
    ws1.cell(row=r_idx, column=1, value="ИТОГО ПО ВСЕМ МОДЕЛЯМ").alignment = align_left
    ws1.cell(row=r_idx, column=2, value=f"=SUM(B9:B{r_idx-1})").number_format = "#,##0"
    ws1.cell(row=r_idx, column=3, value=1.0).number_format = "0.0%"
    ws1.cell(row=r_idx, column=4, value=f"=SUM(D9:D{r_idx-1})").number_format = "#,##0.00 ₽"
    ws1.cell(row=r_idx, column=5, value=f"=SUM(E9:E{r_idx-1})").number_format = "#,##0.00 ₽"
    ws1.cell(row=r_idx, column=6, value=f"=D{r_idx}/B{r_idx}").number_format = "#,##0.00 ₽"
    ws1.cell(row=r_idx, column=7, value=f"=E{r_idx}/B{r_idx}").number_format = "#,##0.00 ₽"
    for col in range(1, 8):
        c = ws1.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if col in [2, 3]: c.alignment = align_center
        elif col >= 4: c.alignment = align_right

    # Table 2: Geography breakdown
    r_idx += 3
    ws1.cell(row=r_idx, column=1, value="2. География продаж GAC по городам РФ").font = f_subhead
    r_idx += 1
    headers_t2 = ["Город / Регион", "Сделок (шт.)", "Доля (%)", "Сумма продаж (₽)", "Комиссия (₽)"]
    for c_i, h in enumerate(headers_t2, 1):
        cell = ws1.cell(row=r_idx, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    city_agg = defaultdict(lambda: {'count': 0, 'price': 0.0, 'comm': 0.0})
    for d in deals:
        c_norm = d['city'].title() if d['city'] else 'Не указан'
        if c_norm.upper() == '2026': c_norm = 'Не указан'
        city_agg[c_norm]['count'] += 1
        city_agg[c_norm]['price'] += d['price']
        city_agg[c_norm]['comm'] += d['comm']

    r_start_c = r_idx + 1
    r_idx += 1
    for c_name, inf in sorted(city_agg.items(), key=lambda x: x[1]['count'], reverse=True):
        ws1.cell(row=r_idx, column=1, value=c_name).alignment = align_left
        ws1.cell(row=r_idx, column=2, value=inf['count']).number_format = "#,##0"
        ws1.cell(row=r_idx, column=3, value=inf['count'] / tot_cnt).number_format = "0.0%"
        ws1.cell(row=r_idx, column=4, value=inf['price']).number_format = "#,##0.00 ₽"
        ws1.cell(row=r_idx, column=5, value=inf['comm']).number_format = "#,##0.00 ₽"
        for col in range(1, 6):
            c = ws1.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if col in [2, 3]: c.alignment = align_center
            elif col >= 4: c.alignment = align_right
        r_idx += 1

    # Total row Table 2
    ws1.cell(row=r_idx, column=1, value="ИТОГО ПО ГОРОДАМ").alignment = align_left
    ws1.cell(row=r_idx, column=2, value=f"=SUM(B{r_start_c}:B{r_idx-1})").number_format = "#,##0"
    ws1.cell(row=r_idx, column=3, value=1.0).number_format = "0.0%"
    ws1.cell(row=r_idx, column=4, value=f"=SUM(D{r_start_c}:D{r_idx-1})").number_format = "#,##0.00 ₽"
    ws1.cell(row=r_idx, column=5, value=f"=SUM(E{r_start_c}:E{r_idx-1})").number_format = "#,##0.00 ₽"
    for col in range(1, 6):
        c = ws1.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if col in [2, 3]: c.alignment = align_center
        elif col >= 4: c.alignment = align_right

    # ==========================================
    # SHEET 2: Москва_Дилеры_Модели
    # ==========================================
    ws2 = wb.create_sheet(title="Москва_Дилеры_Модели")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:I1")
    ws2["A1"] = "ПРОДАЖИ GAC В МОСКВЕ ПО ДИЛЕРАМ В РАЗРЕЗЕ МОДЕЛЕЙ"
    ws2["A1"].font = f_title
    ws2["A1"].alignment = align_left

    ws2.merge_cells("A2:I2")
    ws2["A2"] = "Август 2026 года | Детализация по официальным дилерским холдингам и модельному ряду"
    ws2["A2"].font = f_subtitle
    ws2["A2"].alignment = align_left

    # Section 1: Strictly Moscow
    ws2["A4"] = "РАЗДЕЛ 1. СТРОГО ГОРОД МОСКВА (Город сделки = Москва)"
    ws2["A4"].font = f_subhead

    headers_m = ["Дилер / Холдинг", "Юрлицо / ДЦ в CRM", "GS8", "GS4", "M8 (S9)", "GS3 (S7)", "Hyptec HT", "ИТОГО", "Доля", "Выручка (₽)", "Комиссия (₽)", "Ответственный КАМ"]
    for c_i, h in enumerate(headers_m, 1):
        cell = ws2.cell(row=5, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    # Aggregate Moscow strictly
    m_agg = defaultdict(lambda: {'models': Counter(), 'company': set(), 'price': 0.0, 'comm': 0.0, 'kams': set()})
    for d in moscow_strictly:
        p = d['partner']
        m = d['model_norm']
        m_agg[p]['models'][m] += 1
        m_agg[p]['company'].add(d['company'])
        m_agg[p]['price'] += d['price']
        m_agg[p]['comm'] += d['comm']
        if d['kam']: m_agg[p]['kams'].add(d['kam'])

    r_idx = 6
    r_start_m = r_idx
    tot_moscow_cnt = len(moscow_strictly)
    for p, inf in sorted(m_agg.items(), key=lambda x: sum(x[1]['models'].values()), reverse=True):
        m_cnts = inf['models']
        p_tot = sum(m_cnts.values())
        c_name = ', '.join(sorted(inf['company']))
        kam_str = ', '.join(sorted(inf['kams'])) or '—'

        ws2.cell(row=r_idx, column=1, value=p).alignment = align_left
        ws2.cell(row=r_idx, column=2, value=c_name).alignment = align_left
        ws2.cell(row=r_idx, column=3, value=m_cnts.get('GAC GS8', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=4, value=m_cnts.get('GAC GS4', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=5, value=m_cnts.get('GAC M8 (S9)', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=6, value=m_cnts.get('GAC GS3 (S7)', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=7, value=m_cnts.get('GAC Hyptec HT', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=8, value=p_tot).number_format = "#,##0"
        ws2.cell(row=r_idx, column=9, value=p_tot / tot_moscow_cnt).number_format = "0.0%"
        ws2.cell(row=r_idx, column=10, value=inf['price']).number_format = "#,##0.00 ₽"
        ws2.cell(row=r_idx, column=11, value=inf['comm']).number_format = "#,##0.00 ₽"
        ws2.cell(row=r_idx, column=12, value=kam_str).alignment = align_left

        for col in range(1, 13):
            c = ws2.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if 3 <= col <= 9: c.alignment = align_center
            elif col in [10, 11]: c.alignment = align_right
        r_idx += 1

    # Total row Section 1
    ws2.cell(row=r_idx, column=1, value="ИТОГО ПО МОСКВЕ").alignment = align_left
    ws2.cell(row=r_idx, column=2, value="").alignment = align_left
    for col_c, col_letter in enumerate(['C', 'D', 'E', 'F', 'G', 'H'], start=3):
        ws2.cell(row=r_idx, column=col_c, value=f"=SUM({col_letter}{r_start_m}:{col_letter}{r_idx-1})").number_format = "#,##0"
    ws2.cell(row=r_idx, column=9, value=1.0).number_format = "0.0%"
    ws2.cell(row=r_idx, column=10, value=f"=SUM(J{r_start_m}:J{r_idx-1})").number_format = "#,##0.00 ₽"
    ws2.cell(row=r_idx, column=11, value=f"=SUM(K{r_start_m}:K{r_idx-1})").number_format = "#,##0.00 ₽"
    ws2.cell(row=r_idx, column=12, value="").alignment = align_left
    for col in range(1, 13):
        c = ws2.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if 3 <= col <= 9: c.alignment = align_center
        elif col in [10, 11]: c.alignment = align_right

    # Section 2: Moscow + Moscow Region
    r_idx += 3
    ws2.cell(row=r_idx, column=1, value="РАЗДЕЛ 2. МОСКОВСКИЙ РЕГИОН (Москва + ДЦ в МО: Одинцово, Балашиха, Химки, Нагорное)").font = f_subhead
    r_idx += 1

    headers_mo = ["Дилер / Холдинг", "Юрлицо / Локации ДЦ", "GS8", "GS4", "M8 (S9)", "GS3 (S7)", "Hyptec HT", "ИТОГО", "Доля", "Выручка (₽)", "Комиссия (₽)", "Ответственный КАМ"]
    for c_i, h in enumerate(headers_mo, 1):
        cell = ws2.cell(row=r_idx, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    r_idx += 1
    r_start_mo = r_idx

    mo_agg = defaultdict(lambda: {'models': Counter(), 'cities': Counter(), 'price': 0.0, 'comm': 0.0, 'kams': set()})
    for d in moscow_mo:
        p = d['partner']
        m = d['model_norm']
        mo_agg[p]['models'][m] += 1
        mo_agg[p]['cities'][d['city'].title()] += 1
        mo_agg[p]['price'] += d['price']
        mo_agg[p]['comm'] += d['comm']
        if d['kam']: mo_agg[p]['kams'].add(d['kam'])

    tot_mo_cnt = len(moscow_mo)
    for p, inf in sorted(mo_agg.items(), key=lambda x: sum(x[1]['models'].values()), reverse=True):
        m_cnts = inf['models']
        p_tot = sum(m_cnts.values())
        loc_str = ', '.join([f"{c} ({cnt})" for c, cnt in inf['cities'].most_common()])
        kam_str = ', '.join(sorted(inf['kams'])) or '—'

        ws2.cell(row=r_idx, column=1, value=p).alignment = align_left
        ws2.cell(row=r_idx, column=2, value=loc_str).alignment = align_left
        ws2.cell(row=r_idx, column=3, value=m_cnts.get('GAC GS8', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=4, value=m_cnts.get('GAC GS4', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=5, value=m_cnts.get('GAC M8 (S9)', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=6, value=m_cnts.get('GAC GS3 (S7)', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=7, value=m_cnts.get('GAC Hyptec HT', 0)).number_format = "#,##0"
        ws2.cell(row=r_idx, column=8, value=p_tot).number_format = "#,##0"
        ws2.cell(row=r_idx, column=9, value=p_tot / tot_mo_cnt).number_format = "0.0%"
        ws2.cell(row=r_idx, column=10, value=inf['price']).number_format = "#,##0.00 ₽"
        ws2.cell(row=r_idx, column=11, value=inf['comm']).number_format = "#,##0.00 ₽"
        ws2.cell(row=r_idx, column=12, value=kam_str).alignment = align_left

        for col in range(1, 13):
            c = ws2.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if 3 <= col <= 9: c.alignment = align_center
            elif col in [10, 11]: c.alignment = align_right
        r_idx += 1

    # Total row Section 2
    ws2.cell(row=r_idx, column=1, value="ИТОГО ПО МОСКВЕ И МО").alignment = align_left
    ws2.cell(row=r_idx, column=2, value="").alignment = align_left
    for col_c, col_letter in enumerate(['C', 'D', 'E', 'F', 'G', 'H'], start=3):
        ws2.cell(row=r_idx, column=col_c, value=f"=SUM({col_letter}{r_start_mo}:{col_letter}{r_idx-1})").number_format = "#,##0"
    ws2.cell(row=r_idx, column=9, value=1.0).number_format = "0.0%"
    ws2.cell(row=r_idx, column=10, value=f"=SUM(J{r_start_mo}:J{r_idx-1})").number_format = "#,##0.00 ₽"
    ws2.cell(row=r_idx, column=11, value=f"=SUM(K{r_start_mo}:K{r_idx-1})").number_format = "#,##0.00 ₽"
    ws2.cell(row=r_idx, column=12, value="").alignment = align_left
    for col in range(1, 13):
        c = ws2.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if 3 <= col <= 9: c.alignment = align_center
        elif col in [10, 11]: c.alignment = align_right

    # ==========================================
    # SHEET 3: Реестр_Москва_и_МО
    # ==========================================
    ws3 = wb.create_sheet(title="Реестр_Москва_и_МО")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:N1")
    ws3["A1"] = "РЕЕСТР СДЕЛОК GAC В МОСКВЕ И МО (АВГУСТ 2026)"
    ws3["A1"].font = f_title

    headers_reg = ["№", "ID сделки", "Дата сделки", "Город", "Партнер / Холдинг", "Юрлицо ДЦ", "Бренд", "Модель (норм)", "Модель в CRM", "VIN", "Тип сделки (B2C)", "Цена ТС (₽)", "Комиссия СА (₽)", "Менеджер", "КАМ"]
    for c_i, h in enumerate(headers_reg, 1):
        cell = ws3.cell(row=3, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    r_idx = 4
    for i, d in enumerate(sorted(moscow_mo, key=lambda x: (x['city'], x['partner'], x['deal_id'])), 1):
        ws3.cell(row=r_idx, column=1, value=i).alignment = align_center
        ws3.cell(row=r_idx, column=2, value=int(d['deal_id']) if d['deal_id'].isdigit() else d['deal_id']).alignment = align_center
        ws3.cell(row=r_idx, column=3, value=d['date_str']).alignment = align_center
        ws3.cell(row=r_idx, column=4, value=d['city']).alignment = align_left
        ws3.cell(row=r_idx, column=5, value=d['partner']).alignment = align_left
        ws3.cell(row=r_idx, column=6, value=d['company']).alignment = align_left
        ws3.cell(row=r_idx, column=7, value=d['brand']).alignment = align_center
        ws3.cell(row=r_idx, column=8, value=d['model_norm']).alignment = align_left
        ws3.cell(row=r_idx, column=9, value=d['model_raw']).alignment = align_left
        ws3.cell(row=r_idx, column=10, value=d['vin']).alignment = align_center
        ws3.cell(row=r_idx, column=11, value=d['b2c']).alignment = align_center
        ws3.cell(row=r_idx, column=12, value=d['price']).number_format = "#,##0.00 ₽"
        ws3.cell(row=r_idx, column=13, value=d['comm']).number_format = "#,##0.00 ₽"
        ws3.cell(row=r_idx, column=14, value=d['manager']).alignment = align_left
        ws3.cell(row=r_idx, column=15, value=d['kam']).alignment = align_left

        for col in range(1, 16):
            c = ws3.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if col in [12, 13]: c.alignment = align_right

        r_idx += 1

    # Total row registry Moscow
    ws3.cell(row=r_idx, column=1, value="ИТОГО").alignment = align_center
    ws3.cell(row=r_idx, column=2, value=f"{len(moscow_mo)} сделок").alignment = align_center
    for col in range(3, 12):
        ws3.cell(row=r_idx, column=col, value="").alignment = align_center
    ws3.cell(row=r_idx, column=12, value=f"=SUM(L4:L{r_idx-1})").number_format = "#,##0.00 ₽"
    ws3.cell(row=r_idx, column=13, value=f"=SUM(M4:M{r_idx-1})").number_format = "#,##0.00 ₽"
    for col in range(1, 16):
        c = ws3.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if col in [12, 13]: c.alignment = align_right

    # ==========================================
    # SHEET 4: Все_Сделки_GAC_РФ
    # ==========================================
    ws4 = wb.create_sheet(title="Все_Сделки_GAC_РФ")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:N1")
    ws4["A1"] = "ПОЛНЫЙ РЕЕСТР ВСЕХ СДЕЛОК GAC В РФ (АВГУСТ 2026)"
    ws4["A1"].font = f_title

    for c_i, h in enumerate(headers_reg, 1):
        cell = ws4.cell(row=3, column=c_i, value=h)
        cell.font = f_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    r_idx = 4
    for i, d in enumerate(sorted(deals, key=lambda x: (x['city'], x['partner'], x['deal_id'])), 1):
        ws4.cell(row=r_idx, column=1, value=i).alignment = align_center
        ws4.cell(row=r_idx, column=2, value=int(d['deal_id']) if d['deal_id'].isdigit() else d['deal_id']).alignment = align_center
        ws4.cell(row=r_idx, column=3, value=d['date_str']).alignment = align_center
        ws4.cell(row=r_idx, column=4, value=d['city']).alignment = align_left
        ws4.cell(row=r_idx, column=5, value=d['partner']).alignment = align_left
        ws4.cell(row=r_idx, column=6, value=d['company']).alignment = align_left
        ws4.cell(row=r_idx, column=7, value=d['brand']).alignment = align_center
        ws4.cell(row=r_idx, column=8, value=d['model_norm']).alignment = align_left
        ws4.cell(row=r_idx, column=9, value=d['model_raw']).alignment = align_left
        ws4.cell(row=r_idx, column=10, value=d['vin']).alignment = align_center
        ws4.cell(row=r_idx, column=11, value=d['b2c']).alignment = align_center
        ws4.cell(row=r_idx, column=12, value=d['price']).number_format = "#,##0.00 ₽"
        ws4.cell(row=r_idx, column=13, value=d['comm']).number_format = "#,##0.00 ₽"
        ws4.cell(row=r_idx, column=14, value=d['manager']).alignment = align_left
        ws4.cell(row=r_idx, column=15, value=d['kam']).alignment = align_left

        for col in range(1, 16):
            c = ws4.cell(row=r_idx, column=col)
            c.font = f_regular
            c.border = thin_border
            if col in [12, 13]: c.alignment = align_right

        r_idx += 1

    # Total row registry Russia
    ws4.cell(row=r_idx, column=1, value="ИТОГО ПО РФ").alignment = align_center
    ws4.cell(row=r_idx, column=2, value=f"{len(deals)} сделок").alignment = align_center
    for col in range(3, 12):
        ws4.cell(row=r_idx, column=col, value="").alignment = align_center
    ws4.cell(row=r_idx, column=12, value=f"=SUM(L4:L{r_idx-1})").number_format = "#,##0.00 ₽"
    ws4.cell(row=r_idx, column=13, value=f"=SUM(M4:M{r_idx-1})").number_format = "#,##0.00 ₽"
    for col in range(1, 16):
        c = ws4.cell(row=r_idx, column=col)
        c.font = f_total
        c.fill = fill_total
        c.border = total_border
        if col in [12, 13]: c.alignment = align_right

    # Auto-fit column widths across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                # Skip title lines
                if cell.row in [1, 2]: continue
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save workbook
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"[+] Успешно сохранен в: {OUTPUT_PATH}")

    try:
        wb.save(DESKTOP_PATH)
        print(f"[+] Копия сохранена на рабочий стол: {DESKTOP_PATH}")
    except Exception as e:
        print(f"[!] Не удалось скопировать на рабочий стол: {e}")

if __name__ == '__main__':
    main()
