import json
import re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paths
calc_json_path = '/Users/vaceslavgamaunov/Desktop/Dashboard_usualy/autocalculator/WebAdmin_System/project_data.json'
site_data_path = 'site/data.json'
output_desktop = '/Users/vaceslavgamaunov/Desktop/Сделки_Q3_Витрина_СберАвто_Калькулятор.xlsx'
output_downloads = '/Users/vaceslavgamaunov/Downloads/Сделки_Q3_Витрина_СберАвто_Калькулятор.xlsx'

with open(calc_json_path, 'r', encoding='utf-8') as f:
    calc_data = json.load(f)
calc_cars = calc_data.get('cars', [])

with open(site_data_path, 'r', encoding='utf-8') as f:
    site_data = json.load(f)
sys_db = site_data.get('sys_db', [])

def normalize_brand_model(brand_raw, model_raw, extra_text=''):
    b = str(brand_raw or '').strip().upper()
    m = str(model_raw or '').strip().upper()
    ext = str(extra_text or '').strip().upper()
    full = f"{b} {m} {ext}"
    clean = re.sub(r'[^A-ZА-Я0-9+]', '', full)

    # JETOUR
    if 'JETOUR' in b:
        if 'DASH' in full: return ('Jetour', 'Dashing')
        if any(x in clean for x in ['X70', 'Х70', '70PLUS', '70ПЛЮС', 'X7PLUS', '70+']): return ('Jetour', 'X70 PLUS')
        if any(x in clean for x in ['X90', 'Х90', '90PLUS', '90+']): return ('Jetour', 'X90 PLUS')
        if any(x in clean for x in ['T1', 'Т1', 'TI']): return ('Jetour', 'T1')
        if any(x in clean for x in ['T2', 'Т2']): return ('Jetour', 'T2')
        if 'SOUEAST' in full: return ('Soueast', 'S06')
        return ('Jetour', model_raw.strip())

    # LADA
    if 'LADA' in b or 'ВАЗ' in b:
        if any(x in full for x in ['GRANT', 'GANT', 'GRANR']) or '219' in clean: return ('LADA', 'Granta')
        if 'VEST' in full or 'GFL' in full: return ('LADA', 'Vesta')
        if 'TRAVEL' in full or 'ТРЕВЕЛ' in full or 'ТРЕВ' in full: return ('LADA', 'Niva Travel')
        if any(x in full for x in ['LEGEND', 'BRONTO', 'NIVA', 'НИВА', '4X4']): return ('LADA', 'Niva Legend')
        if 'LARG' in full: return ('LADA', 'Largus')
        if 'ISKR' in full: return ('LADA', 'Iskra')
        if 'AURA' in full: return ('LADA', 'Aura')
        return ('LADA', model_raw.strip())

    # TENET / CHERY & TENET
    if 'TENET' in full or 'ТЕНЕТ' in full:
        if 'TIGGO9' in clean or 'ТИГГО9' in clean: return ('Chery', 'Tiggo 9')
        if 'ARRIZO' in clean: return ('Chery', 'Arrizo 8')
        if 'T4L' in clean or 'Т4L' in clean or 'Т4Л' in clean: return ('Tenet', 'T4L')
        if 'A8' in clean or 'А8' in clean: return ('Tenet', 'A8')
        if 'L4' in clean or 'Л4' in clean: return ('Tenet Plus', 'L4')
        if 'L6' in clean or 'Л6' in clean: return ('Tenet Plus', 'L6')
        if re.search(r'\b[TТ]4\b', full) or 'T4' in m or 'Т4' in m: return ('Tenet', 'T4')
        if re.search(r'\b[TТ]7\b', full) or 'T7' in m or 'Т7' in m: return ('Tenet', 'T7')
        if re.search(r'\b[TТ]8\b', full) or 'T8' in m or 'Т8' in m: return ('Tenet', 'T8')
        if re.search(r'\b[TТ]9\b', full) or 'T9' in m or 'Т9' in m: return ('Tenet', 'T9')
        return ('Tenet', model_raw.strip())

    # SOUEAST
    if 'SOUEAST' in b or 'SOUEAS' in b or 'САУИСТ' in b:
        if '06' in clean: return ('Soueast', 'S06')
        if '07' in clean: return ('Soueast', 'S07')
        if '09' in clean: return ('Soueast', 'S09')
        return ('Soueast', 'S06')

    # GAC
    if 'GAC' in b or 'ГАК' in b:
        if any(x in clean for x in ['GS8', 'CS8']): return ('GAC', 'GS8')
        if any(x in clean for x in ['GS4', 'CS4']): return ('GAC', 'GS4')
        if 'GS3' in clean: return ('GAC', 'GS3')
        if 'EMPOW' in clean: return ('GAC', 'Empow')
        if 'AION' in clean: return ('GAC', 'Aion V')
        if 'HYPTEC' in clean: return ('GAC', 'Hyptec HT')
        if 'M8' in clean or 'М8' in clean: return ('GAC', 'M8')
        if 'S7' in clean: return ('GAC', 'S7')
        if 'S9' in clean or 'GS9' in clean: return ('GAC', 'S9')
        if 'GL' in clean: return ('GAC', 'GS8')
        return ('GAC', model_raw.strip())

    # SOLARIS
    if 'SOLARIS' in b or 'СОЛЯРИС' in b:
        if any(x in clean for x in ['HC', 'НС']): return ('Solaris', 'HC')
        if 'KRX' in clean: return ('Solaris', 'KRX')
        if 'KRS' in clean: return ('Solaris', 'KRS')
        return ('Solaris', 'HC')

    # BELGEE
    if 'BELGEE' in b or 'БЕЛДЖИ' in b:
        if any(x in clean for x in ['50+', 'X50+']): return ('Belgee', 'X50+')
        if any(x in clean for x in ['50', 'X50']) and 'S50' not in clean: return ('Belgee', 'X50')
        if any(x in clean for x in ['70', 'X70']): return ('Belgee', 'X70')
        if 'S50' in clean: return ('Belgee', 'S50')
        return ('Belgee', model_raw.strip())

    # GEELY
    if 'GEELY' in b or 'ДЖИЛИ' in b:
        if 'MONJARO' in clean or 'МАНДЖАРО' in clean: return ('Geely', 'Monjaro')
        if 'CITYRAY' in clean: return ('Geely', 'Cityray')
        if 'COOLRAY' in clean: return ('Geely', 'Coolray')
        if 'ATLAS' in clean: return ('Geely', 'Atlas')
        if 'PREFACE' in clean: return ('Geely', 'Preface')
        if 'OKAVANGO' in clean: return ('Geely', 'Okavango')
        if any(x in clean for x in ['EX5', 'EX5']): return ('Geely', 'EX5')
        if 'S50' in clean: return ('Belgee', 'S50')
        if any(x in clean for x in ['KS001', 'KNEWSTAR']): return ('Knewstar', '001')
        return ('Geely', model_raw.strip())

    # CHANGAN
    if 'CHANGAN' in b or 'ЧАНГАН' in b:
        if any(x in clean for x in ['UNIS', 'CS55']): return ('Changan', 'UNI-S')
        if 'UNIV' in clean: return ('Changan', 'UNI-V')
        if 'UNIK' in clean: return ('Changan', 'UNI-K')
        if 'UNIT' in clean: return ('Changan', 'UNI-T')
        if '35' in clean: return ('Changan', 'CS35PLUS')
        if '75' in clean: return ('Changan', 'CS75PLUS')
        if '95' in clean: return ('Changan', 'CS95')
        if 'ALSVIN' in clean or 'АЛСВИН' in clean: return ('Changan', 'Alsvin')
        if 'EADO' in clean or 'ИДО' in clean: return ('Changan', 'Eado Plus')
        if 'HUNTER' in clean or 'ХАНТЕР' in clean: return ('Changan', 'Hunter Plus')
        if 'LAMORE' in clean or 'ЛАМОР' in clean: return ('Changan', 'Lamore')
        if 'DEEPAL' in clean or 'G318' in clean: return ('Deepal', 'G318')
        return ('Changan', model_raw.strip())

    # HAVAL
    if 'HAVAL' in b or 'ХАВЕЙЛ' in b or 'ХАВАЛ' in b:
        if 'JOLION' in clean or 'ДЖОЛИОН' in clean: return ('Haval', 'Jolion')
        if 'M6' in clean or 'М6' in clean: return ('Haval', 'M6')
        if 'F7X' in clean: return ('Haval', 'F7x')
        if 'F7' in clean: return ('Haval', 'F7')
        if 'DARGO' in clean or 'ДАРГО' in clean: return ('Haval', 'Dargo')
        if 'H3' in clean or 'Н3' in clean: return ('Haval', 'H3')
        if 'H5' in clean or 'Н5' in clean: return ('Haval', 'H5')
        if 'H7' in clean or 'Н7' in clean: return ('Haval', 'H7')
        if 'H9' in clean or 'Н9' in clean: return ('Haval', 'H9')
        if 'POER' in clean: return ('Haval', 'Poer')
        return ('Haval', model_raw.strip())

    # JELAND / JAECOO / OMODA
    if any(k in b for k in ['OMODA', 'JAECOO', 'JELAND', 'JAECO0']):
        if 'C5' in clean or 'С5' in clean or 'CS' in clean: return ('Omoda', 'C5')
        if 'C7' in clean or 'С7' in clean: return ('Omoda', 'C7')
        if 'S5' in clean or 'С5' in clean: return ('Omoda', 'S5')
        if 'J6' in clean:
            return ('Jeland', 'J6') if 'JELAND' in b else ('Jaecoo', 'J6')
        if 'J7' in clean:
            return ('Jeland', 'J7') if 'JELAND' in b else ('Jaecoo', 'J7')
        if 'J8' in clean: return ('Jaecoo', 'J8')
        return (brand_raw.title(), model_raw.strip())

    # CHERY
    if 'CHERY' in b or 'ЧЕРИ' in b:
        if 'TIGGO4' in clean or 'ТИГГО4' in clean: return ('Chery', 'Tiggo 4 Pro')
        if 'TIGGO7' in clean or 'ТИГГО7' in clean: return ('Chery', 'Tiggo 7 Pro Max')
        if 'TIGGO8' in clean or 'ТИГГО8' in clean: return ('Chery', 'Tiggo 8 Pro Max')
        if 'TIGGO9' in clean or 'ТИГГО9' in clean: return ('Chery', 'Tiggo 9')
        if 'ARRIZO' in clean: return ('Chery', 'Arrizo 8')
        return ('Chery', model_raw.strip())

    # X-CITE
    if 'X-CITE' in b or 'XCITE' in b:
        if '7' in clean: return ('X-Cite', 'X-Cross 7')
        if '8' in clean: return ('X-Cite', 'X-Cross 8')
        return ('X-Cite', model_raw.strip())

    # МОСКВИЧ
    if 'МОСКВИЧ' in b or 'MOSKVICH' in b:
        if '3' in clean: return ('Москвич', '3')
        if '8' in clean: return ('Москвич', '8')
        if '70' in clean: return ('Москвич', 'М70')
        if '90' in clean: return ('Москвич', 'М90')
        return ('Москвич', model_raw.strip())

    # KNEWSTAR
    if 'KNEWSTAR' in b: return ('Knewstar', '001')

    # DEEPAL
    if 'DEEPAL' in b: return ('Deepal', 'G318')

    # Параллельный импорт / Другие
    if 'TOYOTA' in b or 'CAMRY' in clean: return ('Toyota', 'Camry')
    if 'MAZDA' in b or 'CX5' in clean: return ('Mazda', 'CX-5')
    if 'HYUNDAI' in b or 'SANTAFE' in clean: return ('Hyundai', 'Santa Fe')
    if 'TANK' in b:
        if '300' in clean: return ('Tank', '300')
        if '500' in clean: return ('Tank', '500')
        if '700' in clean: return ('Tank', '700')
        return ('Tank', model_raw.strip())

    clean_b = brand_raw.strip().title() if brand_raw.isupper() else brand_raw.strip()
    clean_m = model_raw.strip().title() if model_raw.isupper() else model_raw.strip()
    return (clean_b, clean_m)

# 1. Collect all distinct models and their minimum prices from the Calculator project
calc_models = set()
calc_prices = defaultdict(list)

for c in calc_cars:
    b = c.get('Бренд') or c.get('Марка')
    m = c.get('Модель')
    full = c.get('Полное_Имя') or ''
    if not b or 'ПРИОРИТЕТ' in str(m).upper() or 'ПРИОРИТЕТ' in str(b).upper():
        continue

    p_final = c.get('Итоговая цена') or 0
    p_rrc = c.get('РРЦ') or 0
    p_skidka = c.get('Скидка от РРЦ') or 0

    if b == 'LADA' and m == 'NIVA':
        if any(x in full.upper() for x in ['КОМФОРТ', 'КЛАССИК\'24', 'ЛЮКС\'24', 'КХЛ\'24', '1.8']):
            pair = ('LADA', 'Niva Travel')
        else:
            pair = ('LADA', 'Niva Legend')
    else:
        pair = normalize_brand_model(b, m, full)

    calc_models.add(pair)

    if p_rrc == 2026 and p_skidka > 1000000:
        price = p_skidka
    else:
        price = p_final if p_final > 0 else p_rrc

    if price > 50000:
        calc_prices[pair].append(price)

# Explicit base showcase models
calc_models.add(('Toyota', 'Camry'))
calc_models.add(('Mazda', 'CX-5'))
calc_models.add(('Hyundai', 'Santa Fe'))
calc_models.add(('LADA', 'Niva Travel'))
calc_models.add(('LADA', 'Niva Legend'))
calc_models.add(('Tenet', 'T4'))
calc_models.add(('Tenet', 'T8'))
calc_models.add(('Soueast', 'S09'))

# 2. Process Q3 Deals (July, August, September)
q3_months = {'2026-07', '2026-08', '2026-09'}
model_stats = defaultdict(lambda: {
    'total': 0, 'mp': 0, 'retail': 0, 'revenue': 0.0, 'gmv': 0.0,
    'jul': 0, 'aug': 0, 'sep': 0, 'prices': []
})

retail_channels = {'ФДЦ', 'ФДЦ+ГП', 'Online', 'Передача лида'}

for r in sys_db:
    if r.get('SaleQty', 0) != 1: continue
    m = r.get('SaleMonth')
    if m not in q3_months: continue
    
    b = str(r.get('Brand') or '').strip()
    md = str(r.get('Model') or '').strip()
    vin = str(r.get('VIN') or '').strip()
    rev = r.get('Revenue', 0.0) or 0.0
    price = r.get('Price', 0.0) or 0.0
    grp = r.get('ChartGroup', '')
    b2c = r.get('B2C', '')

    norm_pair = normalize_brand_model(b, md, vin)
    
    is_mp = (grp == 'Partners') or (b2c in ['МП1', 'МП2', 'МП3'])
    is_ret = (grp == 'Новые авто') or (b2c in retail_channels)
    
    st = model_stats[norm_pair]
    st['total'] += 1
    st['revenue'] += rev
    st['gmv'] += price
    if is_mp: st['mp'] += 1
    if is_ret: st['retail'] += 1
    if price > 400000: st['prices'].append(price)

    if m == '2026-07': st['jul'] += 1
    elif m == '2026-08': st['aug'] += 1
    elif m == '2026-09': st['sep'] += 1

# Merge calculator models and actual sales
all_active_models = sorted(list(set(list(calc_models) + list(model_stats.keys()))))
print(f"[*] Итоговая витрина моделей (калькулятор + продажи): {len(all_active_models)}")

table_rows = []
for (b, m) in all_active_models:
    st = model_stats[(b, m)]
    tot = st['total']
    mp = st['mp']
    ret = st['retail']
    rev = st['revenue']
    gmv = st['gmv']
    
    mp_share_val = (mp / tot) if tot > 0 else 0.0
    mp_share_str = f"{mp_share_val * 100:.1f}%" if tot > 0 else "0.0%"
    
    # Determine "Цена от ... рублей"
    c_p = calc_prices.get((b, m), [])
    s_p = st['prices']
    if c_p:
        min_price = min(c_p)
    elif s_p:
        min_price = min(s_p)
    else:
        min_price = 0
        
    bm_str = f"{b} {m}".strip()
    table_rows.append({
        'brand_model': bm_str,
        'brand': b,
        'model': m,
        'min_price': int(min_price),
        'q3_deals': tot,
        'mp_deals': mp,
        'retail_deals': ret,
        'mp_share_val': mp_share_val,
        'mp_share_str': mp_share_str,
        'revenue': rev,
        'gmv': gmv,
        'jul': st['jul'],
        'aug': st['aug'],
        'sep': st['sep']
    })

# Primary sorting: Q3 Deals desc, then Revenue desc
table_rows.sort(key=lambda x: (x['q3_deals'], x['revenue']), reverse=True)

# 3. Create Excel Workbook
wb = openpyxl.Workbook()

# Sheet 1: Exact requested 5 columns
ws1 = wb.active
ws1.title = "Сводная витрина Q3"
ws1.views.sheetView[0].showGridLines = True

headers1 = ["Марка/Модель", "Цена от ... рублей", "Сделки ку3", "Доля МП", "Выручка за ку3"]
ws1.append(headers1)

for col_idx in range(1, len(headers1) + 1):
    c = ws1.cell(row=1, column=col_idx)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

tot_deals = sum(r['q3_deals'] for r in table_rows)
tot_mp = sum(r['mp_deals'] for r in table_rows)
tot_rev = sum(r['revenue'] for r in table_rows)
tot_mp_share = (tot_mp / tot_deals) if tot_deals else 0.0

for r_idx, r in enumerate(table_rows, 2):
    c1 = ws1.cell(row=r_idx, column=1, value=r['brand_model'])
    c2 = ws1.cell(row=r_idx, column=2, value=r['min_price'] if r['min_price'] > 0 else "")
    c3 = ws1.cell(row=r_idx, column=3, value=r['q3_deals'])
    c4 = ws1.cell(row=r_idx, column=4, value=r['mp_share_val'])
    c5 = ws1.cell(row=r_idx, column=5, value=r['revenue'])

    c1.alignment = Alignment(horizontal="left", vertical="center")
    c2.alignment = Alignment(horizontal="right", vertical="center")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    c4.alignment = Alignment(horizontal="center", vertical="center")
    c5.alignment = Alignment(horizontal="right", vertical="center")

    c1.font = Font(name="Calibri", size=11, bold=(r['q3_deals'] >= 100))
    c2.font = Font(name="Calibri", size=11)
    c3.font = Font(name="Calibri", size=11, bold=(r['q3_deals'] >= 100))
    c4.font = Font(name="Calibri", size=11)
    c5.font = Font(name="Calibri", size=11)

    c2.number_format = '#,##0 "₽"'
    c3.number_format = '#,##0'
    c4.number_format = '0.0%'
    c5.number_format = '#,##0.00 "₽"'

    if r['q3_deals'] == 0:
        fill_color = "F8FAFC"
        for cell in [c1, c2, c3, c4, c5]:
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            cell.font = Font(name="Calibri", size=11, color="94A3B8")
    elif r['q3_deals'] >= 100:
        fill_color = "F0FDF4"
        for cell in [c1, c2, c3, c4, c5]:
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    for cell in [c1, c2, c3, c4, c5]:
        cell.border = thin_border

# Total Row
tot_row = len(table_rows) + 2
ws1.cell(row=tot_row, column=1, value="ИТОГО").font = Font(name="Calibri", size=11, bold=True)
ws1.cell(row=tot_row, column=2, value="—").font = Font(name="Calibri", size=11, bold=True)
ws1.cell(row=tot_row, column=3, value=tot_deals).font = Font(name="Calibri", size=11, bold=True)
ws1.cell(row=tot_row, column=4, value=tot_mp_share).font = Font(name="Calibri", size=11, bold=True)
ws1.cell(row=tot_row, column=5, value=tot_rev).font = Font(name="Calibri", size=11, bold=True)

ws1.cell(row=tot_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
ws1.cell(row=tot_row, column=3).number_format = '#,##0'
ws1.cell(row=tot_row, column=4).number_format = '0.0%'
ws1.cell(row=tot_row, column=5).number_format = '#,##0.00 "₽"'

for col_idx in range(1, 6):
    cell = ws1.cell(row=tot_row, column=col_idx)
    cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    cell.border = Border(top=Side(style='medium', color='1F497D'), bottom=Side(style='double', color='1F497D'))

ws1.column_dimensions['A'].width = 34
ws1.column_dimensions['B'].width = 22
ws1.column_dimensions['C'].width = 16
ws1.column_dimensions['D'].width = 16
ws1.column_dimensions['E'].width = 24

# Sheet 2: Детализация каналов и месяцев
ws2 = wb.create_sheet(title="Детализация каналов и месяцев")
ws2.views.sheetView[0].showGridLines = True

headers2 = [
    "Марка", "Модель", "Цена от ... рублей", "Сделки Q3 (Всего)", "Сделки МП (Партнеры)", "Сделки Розница (B2C)",
    "Доля МП, %", "Июль 2026", "Август 2026", "Сентябрь 2026", "Выручка за Q3 (TR), ₽", "Оборот авто (GMV), ₽"
]
ws2.append(headers2)

for col_idx in range(1, len(headers2) + 1):
    c = ws2.cell(row=1, column=col_idx)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")

for r_idx, r in enumerate(table_rows, 2):
    ws2.append([
        r['brand'], r['model'], r['min_price'] if r['min_price'] > 0 else "",
        r['q3_deals'], r['mp_deals'], r['retail_deals'],
        r['mp_share_val'], r['jul'], r['aug'], r['sep'], r['revenue'], r['gmv']
    ])
    ws2.cell(row=r_idx, column=3).number_format = '#,##0 "₽"'
    ws2.cell(row=r_idx, column=7).number_format = '0.0%'
    ws2.cell(row=r_idx, column=11).number_format = '#,##0.00 "₽"'
    ws2.cell(row=r_idx, column=12).number_format = '#,##0 "₽"'

for ws in [ws2]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 40)

wb.save(output_desktop)
wb.save(output_downloads)
print(f"[+] Готово! Таблица обновлена и сохранена в:\n  - {output_desktop}\n  - {output_downloads}")
