import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('raw_data/OEM СберАвто финал (19).xlsx', data_only=True)

print(f"Total sheets: {len(wb.sheetnames)}\n")

sheet_data = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    # find header 'город' across all cells
    city_col = None
    header_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column + 1, 40)):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if val.lower() == 'город':
                city_col = c
                header_row = r
                break
        if city_col is not None:
            break
            
    if city_col is None:
        sheet_data[sname] = {
            'has_cities': False,
            'reason': 'Вкладка не содержит таблицы дилеров (столбец "Город" отсутствует)'
        }
        continue
        
    cities_seen = {}
    for r in range(header_row + 1, ws.max_row + 1):
        c_val = ws.cell(row=r, column=city_col).value
        d_val = ws.cell(row=r, column=city_col + 2).value if ws.max_column >= city_col + 2 else None
        
        # Check if row is empty or summary
        if c_val is not None and str(c_val).strip():
            c_str = str(c_val).strip()
            # check if it looks like a note or summary
            if c_str.lower().startswith(('итого', 'примечание', 'всего', 'город')):
                continue
            c_key = c_str.lower()
            if c_key not in cities_seen:
                disp = c_str.capitalize() if c_str.isupper() and len(c_str) > 3 else c_str
                if c_key == 'уфа':
                    disp = 'Уфа'
                cities_seen[c_key] = disp
                
    unique_list = sorted(list(cities_seen.values()), key=lambda x: x.lower())
    sheet_data[sname] = {
        'has_cities': True,
        'header_pos': (header_row, city_col),
        'unique_count': len(unique_list),
        'cities': unique_list
    }

for sname, info in sheet_data.items():
    print("=" * 60)
    print(f"ВКЛАДКА: {sname}")
    if not info['has_cities']:
        print(f"  {info['reason']}")
    elif info['unique_count'] == 0:
        print(f"  Таблица дилеров пуста (0 дилеров)")
    else:
        print(f"  Уникальных городов: {info['unique_count']}")
        for i, c in enumerate(info['cities'], 1):
            print(f"  {i}. {c}")
    print()
