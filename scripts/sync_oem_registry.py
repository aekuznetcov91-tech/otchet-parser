import os
import re
import json
import hashlib
import openpyxl

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DATA_DIR = os.path.join(PROJECT_ROOT, 'raw_data')
SITE_DIR = os.path.join(PROJECT_ROOT, 'site')
DATA_DIR = os.path.join(PROJECT_ROOT, 'data')

BRAND_MAP = {
    'CHERY': 'CHERY & TENET',
    'TENET': 'CHERY & TENET',
    'OMODA': 'OMODA & JAECOO',
    'JAECOO': 'OMODA & JAECOO',
    'GEELY': 'Geely & Belgee',
    'BELGEE': 'Geely & Belgee',
    'LADA': 'LADA',
    'HAVAL': 'HAVAL',
    'CHANGAN': 'CHANGAN',
    'JETOUR': 'JETOUR',
    'SOUEAST': 'Soueast',
    'SOLARIS': 'SOLARIS',
    'GAC': 'GAC',
    'МОСКВИЧ': 'Москвич',
    'DEEPAL': 'DEEPAL',
    'РОЛЬФ': 'Рольф Импорт',
    'JETTA': 'JETTA',
    'VOYAH': 'Voyah'
}

def find_latest_oem_file():
    """Find the most recent OEM Excel file in raw_data or project root."""
    search_dirs = [RAW_DATA_DIR, PROJECT_ROOT]
    candidates = []
    
    for sdir in search_dirs:
        if not os.path.exists(sdir):
            continue
        for fname in os.listdir(sdir):
            if fname.startswith('~$') or fname.startswith('.'):
                continue
            if 'oem' in fname.lower() and fname.lower().endswith(('.xlsx', '.xlsm')):
                fpath = os.path.join(sdir, fname)
                mtime = os.path.getmtime(fpath)
                num_match = re.search(r'\((\d+)\)', fname)
                ver_num = int(num_match.group(1)) if num_match else 0
                candidates.append((ver_num, mtime, fpath, fname))
                
    if not candidates:
        return None
        
    candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return candidates[0][2]

def extract_dealers_from_oem(fpath):
    """Extract dealer centers (roofs) across all brand sheets in the OEM file."""
    wb = openpyxl.load_workbook(fpath, data_only=True)
    all_dealers = []
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        header_found = False
        headers = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            vals = [str(c).strip() if c is not None else '' for c in row]
            if not any(vals):
                continue
            row_str = ' '.join(vals).lower()
            if 'город' in row_str and ('инн' in row_str or 'юл' in row_str or 'фдц' in row_str or 'название' in row_str):
                header_found = True
                headers = [v.lower() for v in vals]
                continue
                
            if header_found and any(vals):
                c0 = vals[0] if len(vals) > 0 else ''
                if not c0 or c0.startswith('http') or c0.startswith('=') or c0.lower() == 'город':
                    continue
                
                entry = {'sheet': sname}
                for c_idx, val in enumerate(vals):
                    if c_idx < len(headers):
                        h = headers[c_idx]
                        if 'город' in h:
                            entry['city'] = val
                        elif 'ссылка' in h or 'чат' in h or 'групп' in h:
                            entry['group_link'] = val
                        elif 'название в беке' in h or 'название в бэке' in h or 'бэк' in h or 'бек' in h:
                            entry['back_name'] = val
                        elif 'название' in h or 'дц' in h or 'дилер' in h:
                            if 'name' not in entry:
                                entry['name'] = val
                        elif 'юл' in h or 'юр' in h:
                            entry['legal_entity'] = val
                        elif 'инн' in h:
                            clean_inn = str(val).split('.')[0].strip()
                            entry['inn'] = clean_inn if clean_inn != '-' else ''
                        elif 'фдц' in h or 'код' in h:
                            clean_fdc = str(val).split('.')[0].strip()
                            entry['fdc_code'] = clean_fdc if clean_fdc != '-' else ''
                        elif 'адрес' in h:
                            entry['address'] = val
                        elif 'почт' in h or 'email' in h:
                            entry['email'] = val
                        elif 'ответствен' in h or 'куратор' in h or 'координатор' in h:
                            entry['responsible'] = val
                
                s_up = sname.upper()
                brand = sname
                for k, v in BRAND_MAP.items():
                    if k in s_up:
                        brand = v
                        break
                entry['brand'] = brand
                all_dealers.append(entry)
                
    wb.close()
    return all_dealers

def sync_oem_to_registry(oem_file_path=None):
    """
    Sync OEM dealers into partners_registry.json, preserving all KAM assignments and aliases.
    Regenerates russia_dealer_benchmarks.json automatically.
    """
    if not oem_file_path:
        oem_file_path = find_latest_oem_file()
        
    if not oem_file_path or not os.path.exists(oem_file_path):
        print('[*] OEM файл не найден для синхронизации.')
        return False
        
    print(f'[*] Синхронизация реестра OEM из файла: {os.path.basename(oem_file_path)}')
    dealers = extract_dealers_from_oem(oem_file_path)
    print(f'[*] Извлечено {len(dealers)} записей ДЦ из OEM файла.')
    
    registry_paths = [
        os.path.join(SITE_DIR, 'partners_registry.json'),
        os.path.join(DATA_DIR, 'partners_registry.json'),
        os.path.join(PROJECT_ROOT, 'partners_registry.json')
    ]
    
    reg_data = None
    target_path = registry_paths[0]
    for rp in registry_paths:
        if os.path.exists(rp):
            try:
                with open(rp, 'r', encoding='utf-8') as f:
                    reg_data = json.load(f)
                target_path = rp
                break
            except Exception:
                pass
                
    if not reg_data or 'partners' not in reg_data:
        reg_data = {'partners': [], 'unmatched_queue': [], 'auto_matched_deals': [], 'auto_matched_leads': []}
        
    partners = reg_data.get('partners', [])
    partner_by_id = {p['partner_id']: p for p in partners}
    
    inn_map = {}
    fdc_map = {}
    name_map = {}
    
    for p in partners:
        cname = p['canonical_name'].strip().lower()
        name_map[cname] = p
        for a in p.get('bitrix_aliases', []) + p.get('bi_aliases', []) + p.get('pochta_aliases', []):
            if a:
                name_map[a.strip().lower()] = p
        for o in p.get('oem_data', []):
            inn = str(o.get('inn', '')).split('.')[0].strip()
            fdc = str(o.get('fdc_code', '')).split('.')[0].strip()
            leg = str(o.get('legal_entity', '')).strip().lower()
            nm = str(o.get('name', '')).strip().lower()
            bnm = str(o.get('back_name', '')).strip().lower()
            if inn: inn_map[inn] = p
            if fdc: fdc_map[fdc] = p
            if leg: name_map[leg] = p
            if nm: name_map[nm] = p
            if bnm: name_map[bnm] = p

    # Clear previous oem_data to cleanly refresh
    for p in partners:
        p['oem_data'] = []
                
    created_partners = 0
    for d in dealers:
        inn = str(d.get('inn', '')).strip()
        fdc = str(d.get('fdc_code', '')).strip()
        leg = str(d.get('legal_entity', '')).strip().lower()
        nm = str(d.get('name', '')).strip().lower()
        bnm = str(d.get('back_name', '')).strip().lower()
        
        p = None
        if inn and inn in inn_map:
            p = inn_map[inn]
        elif fdc and fdc in fdc_map:
            p = fdc_map[fdc]
        elif leg and leg in name_map:
            p = name_map[leg]
        elif nm and nm in name_map:
            p = name_map[nm]
        elif bnm and bnm in name_map:
            p = name_map[bnm]
            
        if not p:
            cname = d.get('name') or d.get('legal_entity') or 'OEM Партнер'
            pid_hash = hashlib.md5(f"{cname}_{inn}_{fdc}".encode('utf-8')).hexdigest()[:8]
            pid = f"P_OEM_{pid_hash}"
            p = {
                'partner_id': pid,
                'canonical_name': cname,
                'kam': 'Не назначен',
                'bitrix_aliases': [],
                'bi_aliases': [],
                'pochta_aliases': [],
                'oem_data': []
            }
            partners.append(p)
            partner_by_id[pid] = p
            created_partners += 1
            if inn: inn_map[inn] = p
            if fdc: fdc_map[fdc] = p
            if leg: name_map[leg] = p
            if nm: name_map[nm] = p
            if bnm: name_map[bnm] = p
            
        # Add dealer if not duplicate in this partner
        is_dup = False
        for ex in p['oem_data']:
            if (ex.get('brand') == d.get('brand') and 
                ex.get('city') == d.get('city') and 
                ex.get('name') == d.get('name') and
                ex.get('address') == d.get('address')):
                is_dup = True
                break
        if not is_dup:
            p['oem_data'].append(d)
            if inn and inn not in inn_map:
                inn_map[inn] = p
            if fdc and fdc not in fdc_map:
                fdc_map[fdc] = p
                
    total_assigned = sum(len(p.get('oem_data', [])) for p in partners)
    print(f"[*] Реестр обновлен: {len(partners)} Master Partners ({created_partners} новых), {total_assigned} привязанных ДЦ OEM.")
    
    # Save to all registry paths
    for rp in registry_paths:
        os.makedirs(os.path.dirname(rp), exist_ok=True)
        with open(rp, 'w', encoding='utf-8') as f:
            json.dump(reg_data, f, ensure_ascii=False, indent=2)

    # Re-apply manual partner splits & KAM overrides
    try:
        from scripts.fix_partners_splits import apply_partners_splits
        apply_partners_splits()
    except Exception as e:
        print(f"[!] Warning applying partner splits: {e}")
            
    # Trigger benchmark update
    try:
        import sys
        if PROJECT_ROOT not in sys.path:
            sys.path.insert(0, PROJECT_ROOT)
        from scripts.fetch_market_dealers import generate_benchmark_data
        generate_benchmark_data()
        print('[*] Бенчмарки дилерских сетей РФ успешно пересчитаны.')
    except Exception as e:
        print(f'[!] Warning updating benchmarks: {e}')
        
    return True

if __name__ == '__main__':
    sync_oem_to_registry()
