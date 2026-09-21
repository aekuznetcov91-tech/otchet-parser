import os
import sys
import re
import json
import openpyxl
from collections import defaultdict

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DATA_DIR = os.path.join(PROJECT_ROOT, 'raw_data')
SITE_DIR = os.path.join(PROJECT_ROOT, 'site')
DATA_DIR = os.path.join(PROJECT_ROOT, 'data')

def clean_inn(val):
    if val is None:
        return ''
    s = re.sub(r'\D', '', str(val).split('.')[0].strip())
    if not s or s == '0':
        return ''
    if len(s) == 9 or len(s) == 11:
        s = '0' + s
    return s

def clean_str(val):
    if val is None:
        return ''
    return re.sub(r'\s+', ' ', str(val)).strip()

def normalize_kam(val):
    if not val:
        return 'Не назначен'
    s = clean_str(val)
    s_low = s.lower()
    if 'кузнецов' in s_low:
        return 'Андрей Кузнецов'
    if 'чихарев' in s_low or 'чихарёв' in s_low:
        return 'Алексей Чихарев'
    if 'дариенко' in s_low:
        return 'Светлана Дариенко'
    if 'солдатова' in s_low:
        return 'Валерия Солдатова'
    if 'добролюбова' in s_low:
        return 'Евгения Добролюбова'
    return s

def run_enrichment():
    print("=" * 60)
    print("[*] ЗАПУСК СКВОЗНОГО МАТЧИНГА: BI + БИТРИКС + OEM (15)")
    print("=" * 60)

    # 1. Load OEM (15)
    oem_file = os.path.join(RAW_DATA_DIR, 'OEM СберАвто финал (15).xlsx')
    if not os.path.exists(oem_file):
        raise FileNotFoundError(f"OEM 15 file not found: {oem_file}")

    print(f"[*] Загрузка OEM дилеров из: {os.path.basename(oem_file)}")
    wb_oem = openpyxl.load_workbook(oem_file, data_only=True)
    oem_by_inn = defaultdict(list)
    oem_kam_by_inn = {}

    for sname in wb_oem.sheetnames:
        ws = wb_oem[sname]
        headers = []
        for r in ws.iter_rows(values_only=True):
            if not any(r): continue
            r_str = ' '.join(str(c) for c in r if c).lower()
            if 'город' in r_str and ('инн' in r_str or 'юл' in r_str or 'дц' in r_str or 'название' in r_str):
                headers = [str(c or '').strip().lower() for c in r]
                continue
            if headers:
                row_dict = dict(zip(headers, r))
                inn_raw = None
                for hk in ['инн', 'инн юл', 'инн компании']:
                    if hk in row_dict and row_dict[hk]:
                        inn_raw = row_dict[hk]
                        break
                inn = clean_inn(inn_raw)
                if not inn: continue

                kam_raw = None
                for hk in ['ответственный', 'ответственный кам', 'кам', 'куратор', 'менеджер']:
                    if hk in row_dict and row_dict[hk]:
                        kam_raw = row_dict[hk]
                        break
                kam = normalize_kam(kam_raw)

                name_raw = None
                for hk in ['название', 'дц', 'название дц', 'дилер']:
                    if hk in row_dict and row_dict[hk]:
                        name_raw = row_dict[hk]
                        break
                name = clean_str(name_raw)

                legal_raw = None
                for hk in ['юридическое лицо', 'юр лицо', 'юрлицо', 'юл']:
                    if hk in row_dict and row_dict[hk]:
                        legal_raw = row_dict[hk]
                        break
                legal = clean_str(legal_raw)

                city_raw = row_dict.get('город', '')
                city = clean_str(city_raw)

                entry = {
                    'sheet': sname,
                    'city': city,
                    'name': name,
                    'legal_entity': legal,
                    'inn': inn,
                    'responsible': kam
                }
                oem_by_inn[inn].append(entry)
                if kam and kam not in ('Не назначен', '', '-'):
                    oem_kam_by_inn[inn] = kam

    wb_oem.close()
    print(f"    - Найдено {len(oem_by_inn)} уникальных ИНН в OEM (15).")
    print(f"    - С привязанным КАМом: {len(oem_kam_by_inn)} ИНН.")

    # 2. Load дилеры с инн и городом.xlsx (BI dealer names)
    bi_file = os.path.join(RAW_DATA_DIR, 'дилеры с инн и городом.xlsx')
    bi_dealers = []
    if os.path.exists(bi_file):
        print(f"[*] Загрузка BI дилеров из: {os.path.basename(bi_file)}")
        wb_bi = openpyxl.load_workbook(bi_file, data_only=True)
        ws_bi = wb_bi.active
        for r in ws_bi.iter_rows(min_row=2, values_only=True):
            dname = clean_str(r[0])
            inn = clean_inn(r[1])
            city = clean_str(r[2])
            if dname and inn:
                bi_dealers.append({'dealer_name': dname, 'inn': inn, 'city': city})
        wb_bi.close()
        print(f"    - Загружено {len(bi_dealers)} записей дилеров BI.")
    else:
        print("[!] Внимание: файл 'дилеры с инн и городом.xlsx' не найден!")

    # 3. Load отчет_сентябрь.xlsx (Bitrix companies with INN and Deal links)
    sep_file = os.path.join(RAW_DATA_DIR, 'отчет_сентябрь.xlsx')
    if not os.path.exists(sep_file):
        raise FileNotFoundError(f"Отчет сентябрь не найден: {sep_file}")

    print(f"[*] Загрузка подтвержденных сделок сентября из: {os.path.basename(sep_file)}")
    wb_sep = openpyxl.load_workbook(sep_file, data_only=True)
    ws_sep = wb_sep.active

    sep_deals_by_id = {}
    sep_deals_by_vin = {}
    sep_companies_by_inn = defaultdict(set)

    for r in ws_sep.iter_rows(min_row=3, values_only=True):
        if not r or not any(r): continue
        cid = clean_str(r[3]) if len(r) > 3 else ''
        cname = clean_str(r[4]) if len(r) > 4 else ''
        inn = clean_inn(r[5]) if len(r) > 5 else ''
        price = float(r[8] or 0) if len(r) > 8 and r[8] else 0.0
        comm = float(r[9] or 0) if len(r) > 9 and r[9] else 0.0
        dkp_num = clean_str(r[10]) if len(r) > 10 else ''
        dkp_date = str(r[11]).split()[0] if len(r) > 11 and r[11] else ''
        brand = clean_str(r[12]) if len(r) > 12 else ''
        model = clean_str(r[13]) if len(r) > 13 else ''
        vin = clean_str(r[15]) if len(r) > 15 else ''
        client_id = clean_str(r[21]) if len(r) > 21 else ''
        city = clean_str(r[22]) if len(r) > 22 else ''
        link = str(r[24] if len(r) > 24 and r[24] else (r[25] if len(r) > 25 else ''))
        deal_id_m = re.search(r'/details/(\d+)/', link)
        deal_id = deal_id_m.group(1) if deal_id_m else ''

        deal_info = {
            'company_id': cid,
            'company_name': cname,
            'inn': inn,
            'price': price,
            'comm': comm,
            'dkp_number': dkp_num,
            'dkp_date': dkp_date,
            'brand': brand,
            'model': model,
            'vin': vin,
            'client_id': client_id,
            'city': city,
            'deal_id': deal_id
        }

        if deal_id:
            sep_deals_by_id[deal_id] = deal_info
        if vin:
            sep_deals_by_vin[vin] = deal_info
        if cname and inn:
            sep_companies_by_inn[inn].add(cname)

    wb_sep.close()
    print(f"    - Загружено {len(sep_deals_by_id)} сделок с DealId и {len(sep_companies_by_inn)} уникальных ИНН.")

    # 4. Load partners_registry.json
    reg_paths = [
        os.path.join(SITE_DIR, 'partners_registry.json'),
        os.path.join(DATA_DIR, 'partners_registry.json'),
        os.path.join(PROJECT_ROOT, 'partners_registry.json')
    ]
    target_reg_path = reg_paths[0]
    reg_data = None
    for rp in reg_paths:
        if os.path.exists(rp):
            with open(rp, 'r', encoding='utf-8') as f:
                reg_data = json.load(f)
            target_reg_path = rp
            break

    if not reg_data or 'partners' not in reg_data:
        raise ValueError("Invalid partners_registry.json")

    partners = reg_data['partners']
    print(f"[*] Текущее число Master Partners в реестре: {len(partners)}")

    # Index partners by INN, bitrix_alias, bi_alias, and canonical_name
    p_by_inn = {}
    p_by_name = {}

    for p in partners:
        # Check OEM data INNs
        for o in p.get('oem_data', []):
            o_inn = clean_inn(o.get('inn'))
            if o_inn: p_by_inn[o_inn] = p

        cname_clean = clean_str(p.get('canonical_name')).lower()
        if cname_clean: p_by_name[cname_clean] = p

        for a in p.get('bitrix_aliases', []):
            a_clean = clean_str(a).lower()
            if a_clean: p_by_name[a_clean] = p

        for a in p.get('bi_aliases', []):
            a_clean = clean_str(a).lower()
            if a_clean: p_by_name[a_clean] = p

    # Inviolable KAM Matrix rules
    def get_matrix_override(cname_str, city_str=""):
        cl = cname_str.lower()
        ct = (city_str or "").lower()

        # Andrey Kuznetsov
        if 'рольф' in cl: return "Андрей Кузнецов"
        if any(k in cl for k in ['агат', 'автопрофиль', 'аркада', 'платинум', 'квант', 'альтаир', 'приоритет моторс', 'максима авто', 'планета авто', 'гольфстрим', 'lucky motors', 'эксперт самара', 'автолидер']):
            return "Андрей Кузнецов"
        if 'олимп' in cl or 'темп авто кубань' in cl: return "Андрей Кузнецов"
        if 'армада-авто' in cl and 'online ооо "тд' in cl: return "Андрей Кузнецов"

        # Aleksei Chikharev
        if 'кунцево' in cl: return "Алексей Чихарев"
        if 'армада-авто' in cl and 'online ооо "тд' not in cl: return "Алексей Чихарев"
        if 'оренбург' in cl or 'эксперт авто оренбург' in cl or 'эксперт св' in cl: return "Алексей Чихарев"
        if 'автоимпорт' in cl or 'маркар' in cl or 'диалог' in cl or 'автодом' in cl: return "Алексей Чихарев"
        if 'спектр' in cl and 'апельсин' in cl: return "Алексей Чихарев"

        # Svetlana Darienko
        if 'максимум' in cl or 'вагнер' in cl or 'автопилот' in cl or 'слово' in cl or 'сибкар' in cl or 'лидер сервис' in cl:
            return "Светлана Дариенко"
        if 'аларм' in cl: return "Светлана Дариенко"
        if 'фаворит' in cl and ('санкт-петербург' in cl or 'спб' in cl): return "Светлана Дариенко"
        if 'автоград' in cl:
            return "Светлана Дариенко" if 'калининград' in ct else "Алексей Чихарев"
        if 'премиум авто' in cl:
            return "Светлана Дариенко" if ('спб' in ct or 'санкт-петербург' in ct) else "Алексей Чихарев"
        if 'авторитэйл м' in cl or 'авторитэйл' in cl:
            if 'санкт-петербург' in ct or 'спб' in ct: return "Светлана Дариенко"
            return "Валерия Солдатова"

        # Valeria Soldatova
        if any(k in cl for k in ['темп авто к', 'темп авто дон', 'техно-темп', 'трансфор', 'фининвест', 'автосфера', 'автоюг', 'артекс', 'ринг']):
            return "Валерия Солдатова"

        return None

    # 5. Enrich with BI dealers (дилеры с инн и городом)
    bi_matched = 0
    for bd in bi_dealers:
        dname = bd['dealer_name']
        inn = bd['inn']
        p = p_by_inn.get(inn)
        if not p:
            p = p_by_name.get(dname.lower())
        if p:
            bi_matched += 1
            if 'bi_aliases' not in p: p['bi_aliases'] = []
            if dname not in p['bi_aliases']:
                p['bi_aliases'].append(dname)
            p_by_inn[inn] = p
            p_by_name[dname.lower()] = p

    print(f"[*] Сметчено BI дилеров: {bi_matched} из {len(bi_dealers)}.")

    # 6. Enrich with September Bitrix companies (отчет_сентябрь.xlsx)
    sep_matched = 0
    new_partners_created = 0

    max_pid = 1000
    for p in partners:
        try:
            pid_int = int(p.get('partner_id', 0))
            if pid_int > max_pid: max_pid = pid_int
        except Exception:
            pass

    for inn, cnames in sep_companies_by_inn.items():
        p = p_by_inn.get(inn)
        if not p:
            for cn in cnames:
                p = p_by_name.get(cn.lower())
                if p: break

        if p:
            sep_matched += 1
            if 'bitrix_aliases' not in p: p['bitrix_aliases'] = []
            for cn in cnames:
                if cn not in p['bitrix_aliases']:
                    p['bitrix_aliases'].append(cn)
                p_by_name[cn.lower()] = p
            p_by_inn[inn] = p

            # Check KAM update: OEM (15) priority with matrix guard
            first_cname = list(cnames)[0]
            mat_override = get_matrix_override(first_cname)
            if mat_override:
                p['kam'] = mat_override
            elif inn in oem_kam_by_inn:
                p['kam'] = oem_kam_by_inn[inn]
        else:
            # Create new partner entry
            max_pid += 1
            first_cname = list(cnames)[0]
            mat_override = get_matrix_override(first_cname)
            assigned_kam = mat_override or oem_kam_by_inn.get(inn, "Не назначен")

            new_p = {
                'partner_id': max_pid,
                'canonical_name': first_cname.replace('ONLINE', '').replace('Online', '').strip(' "\'\t'),
                'kam': assigned_kam,
                'bitrix_aliases': list(cnames),
                'bi_aliases': [],
                'pochta_aliases': [],
                'oem_data': oem_by_inn.get(inn, [])
            }
            partners.append(new_p)
            p_by_inn[inn] = new_p
            for cn in cnames:
                p_by_name[cn.lower()] = new_p
            new_partners_created += 1

    print(f"[*] Сметчено компаний сентября: {sep_matched} из {len(sep_companies_by_inn)}.")
    print(f"[*] Создано новых Master Partners по ИНН: {new_partners_created}.")
    print(f"[*] Итого Master Partners в реестре: {len(partners)}.")

    # 7. Update OEM data and KAMs from OEM (15) across all partners
    oem_updated = 0
    for p in partners:
        cname = p.get('canonical_name', '')
        mat_override = get_matrix_override(cname)

        # Gather all INNs for this partner
        p_inns = set()
        for o in p.get('oem_data', []):
            o_inn = clean_inn(o.get('inn'))
            if o_inn: p_inns.add(o_inn)

        for inn in p_inns:
            if inn in oem_by_inn:
                # Add any missing OEM entries
                existing_descs = set(f"{o.get('sheet')}_{o.get('name')}_{o.get('address')}" for o in p['oem_data'])
                for o_fresh in oem_by_inn[inn]:
                    desc = f"{o_fresh.get('sheet')}_{o_fresh.get('name')}_{o_fresh.get('address')}"
                    if desc not in existing_descs:
                        p['oem_data'].append(o_fresh)
                        existing_descs.add(desc)

                if not mat_override and inn in oem_kam_by_inn:
                    p['kam'] = normalize_kam(oem_kam_by_inn[inn])
                    oem_updated += 1

        if mat_override:
            p['kam'] = mat_override

    print(f"[*] Актуализировано КАМов по OEM (15): {oem_updated}.")

    # Save to all registry paths
    reg_data['partners'] = partners
    reg_data['total_master_partners'] = len(partners)
    for rp in reg_paths:
        os.makedirs(os.path.dirname(rp), exist_ok=True)
        with open(rp, 'w', encoding='utf-8') as f:
            json.dump(reg_data, f, ensure_ascii=False, indent=2)

    # Save September deals bridge file for parser_engine.py
    bridge_file = os.path.join(DATA_DIR, 'september_deals_bridge.json')
    with open(bridge_file, 'w', encoding='utf-8') as f:
        json.dump({
            'deals_by_id': sep_deals_by_id,
            'deals_by_vin': sep_deals_by_vin
        }, f, ensure_ascii=False, indent=2)
    print(f"[*] Сохранен мост сделок сентября: {bridge_file} ({len(sep_deals_by_id)} сделок).")
    print("=" * 60)
    print("✅ СКВОЗНОЙ МАТЧИНГ УСПЕШНО ЗАВЕРШЕН!")
    print("=" * 60)

if __name__ == '__main__':
    run_enrichment()
