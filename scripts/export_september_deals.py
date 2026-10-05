import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict, Counter

def build_report():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    json_path = os.path.join(base_dir, 'site', 'data.json')
    
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    deals = [x for x in data.get('sys_db_partners', []) if x.get('Month') == '2026-09' and x.get('Type') == 'Сделка']
    print(f"Total September deals loaded: {len(deals)}")

    # Aggregate by dealer and brand
    dealer_brand_counts = defaultdict(lambda: defaultdict(int))
    dealer_totals = defaultdict(int)
    brand_totals = Counter()

    for d in deals:
        dealer = d.get('Partner') or d.get('RawPartner') or 'Не указан'
        brand = d.get('Brand') or 'Не указан'
        dealer_brand_counts[dealer][brand] += 1
        dealer_totals[dealer] += 1
        brand_totals[brand] += 1

    sorted_dealers = sorted(dealer_totals.keys(), key=lambda d: (-dealer_totals[d], d))
    sorted_brands = [b for b, _ in brand_totals.most_common()]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_family = 'Segoe UI'
    title_font = Font(name=font_family, size=15, bold=True, color='1F497D')
    sub_font = Font(name=font_family, size=10, italic=True, color='595959')
    header_font = Font(name=font_family, size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    header_matrix_fill = PatternFill(start_color='203764', end_color='203764', fill_type='solid')
    brand_header_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    total_header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    data_font = Font(name=font_family, size=10)
    bold_font = Font(name=font_family, size=10, bold=True)
    total_font = Font(name=font_family, size=11, bold=True)
    total_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    zebra_fill = PatternFill(start_color='F9FAFC', end_color='F9FAFC', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='D9D9D9')
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_top = Side(border_style='thin', color='1F4E78')
    double_bottom = Side(border_style='double', color='1F4E78')
    total_border = Border(top=thick_top, bottom=double_bottom, left=thin_border_side, right=thin_border_side)

    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # SHEET 1: Разбивка по дилерам и брендам
    ws1 = wb.create_sheet(title='Разбивка по дилерам и брендам')
    ws1.views.sheetView[0].showGridLines = True

    ws1['A1'] = 'Отчёт по сделкам: Дилеры и Бренды автомобилей'
    ws1['A1'].font = title_font
    ws1['A2'] = f'Период: Сентябрь 2026 г. | Всего сделок: {len(deals):,} | Всего дилеров: {len(sorted_dealers)} | Брендов: {len(sorted_brands)}'.replace(',', ' ')
    ws1['A2'].font = sub_font

    headers1 = ['№ п/п', 'Дилер / Холдинг', 'Бренд автомобиля', 'Количество сделок']
    start_row1 = 4

    for col_idx, h in enumerate(headers1, 1):
        c = ws1.cell(row=start_row1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_header
        c.border = data_border
    ws1.row_dimensions[start_row1].height = 26

    row_num = start_row1 + 1
    item_idx = 1

    for dealer in sorted_dealers:
        d_brands = sorted(dealer_brand_counts[dealer].items(), key=lambda x: (-x[1], x[0]))
        for brand, count in d_brands:
            ws1.cell(row=row_num, column=1, value=item_idx).alignment = align_center
            ws1.cell(row=row_num, column=2, value=dealer).alignment = align_left
            ws1.cell(row=row_num, column=3, value=brand).alignment = align_left
            ws1.cell(row=row_num, column=4, value=count).alignment = align_right
            ws1.cell(row=row_num, column=4).number_format = '#,##0'
            
            is_even = (item_idx % 2 == 0)
            for col_idx in range(1, 5):
                cell = ws1.cell(row=row_num, column=col_idx)
                cell.font = data_font
                cell.border = data_border
                if is_even:
                    cell.fill = zebra_fill
            
            ws1.row_dimensions[row_num].height = 20
            row_num += 1
            item_idx += 1

    ws1.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    ws1.cell(row=row_num, column=1, value='ИТОГО СДЕЛОК').alignment = Alignment(horizontal='right', vertical='center')
    tot_sum_cell = ws1.cell(row=row_num, column=4, value=f'=SUM(D{start_row1+1}:D{row_num-1})')
    tot_sum_cell.alignment = align_right
    tot_sum_cell.number_format = '#,##0'

    for col_idx in range(1, 5):
        c = ws1.cell(row=row_num, column=col_idx)
        c.font = total_font
        c.fill = total_fill
        c.border = total_border
    ws1.row_dimensions[row_num].height = 24

    ws1.auto_filter.ref = f'A{start_row1}:D{row_num-1}'
    ws1.freeze_panes = f'A{start_row1+1}'

    ws1.column_dimensions['A'].width = 8
    ws1.column_dimensions['B'].width = 44
    ws1.column_dimensions['C'].width = 24
    ws1.column_dimensions['D'].width = 20

    # SHEET 2: Сводная матрица (Дилер x Бренд)
    ws2 = wb.create_sheet(title='Сводная матрица (Дилер x Бренд)')
    ws2.views.sheetView[0].showGridLines = True

    ws2['A1'] = 'Матрица распределения сделок: Дилер x Бренд автомобиля'
    ws2['A1'].font = title_font
    ws2['A2'] = f'Период: Сентябрь 2026 г. | Сделок: {len(deals):,} | Строки — дилеры (по убыванию объёма), Столбцы — бренды'.replace(',', ' ')
    ws2['A2'].font = sub_font

    start_row2 = 4
    matrix_headers = ['№ п/п', 'Дилер / Холдинг'] + sorted_brands + ['ИТОГО сделок']

    for col_idx, h in enumerate(matrix_headers, 1):
        c = ws2.cell(row=start_row2, column=col_idx, value=h)
        c.font = header_font
        if col_idx <= 2:
            c.fill = header_matrix_fill
        elif col_idx == len(matrix_headers):
            c.fill = total_header_fill
        else:
            c.fill = brand_header_fill
        c.alignment = align_header
        c.border = data_border
    ws2.row_dimensions[start_row2].height = 32

    row_num2 = start_row2 + 1
    for d_idx, dealer in enumerate(sorted_dealers, 1):
        ws2.cell(row=row_num2, column=1, value=d_idx).alignment = align_center
        ws2.cell(row=row_num2, column=2, value=dealer).alignment = align_left
        
        first_col_letter = get_column_letter(3)
        last_col_letter = get_column_letter(2 + len(sorted_brands))
        
        for b_idx, brand in enumerate(sorted_brands, 3):
            cnt = dealer_brand_counts[dealer].get(brand, 0)
            c = ws2.cell(row=row_num2, column=b_idx, value=cnt if cnt > 0 else '')
            c.alignment = align_right
            if cnt > 0:
                c.number_format = '#,##0'
        
        tot_c = ws2.cell(row=row_num2, column=len(matrix_headers), value=f'=SUM({first_col_letter}{row_num2}:{last_col_letter}{row_num2})')
        tot_c.alignment = align_right
        tot_c.number_format = '#,##0'
        tot_c.font = bold_font
        
        is_even = (d_idx % 2 == 0)
        for col_idx in range(1, len(matrix_headers) + 1):
            cell = ws2.cell(row=row_num2, column=col_idx)
            if col_idx != len(matrix_headers):
                cell.font = data_font
            cell.border = data_border
            if is_even and col_idx < len(matrix_headers):
                cell.fill = zebra_fill
            elif col_idx == len(matrix_headers):
                cell.fill = PatternFill(start_color='EEF2F9', end_color='EEF2F9', fill_type='solid')
                
        ws2.row_dimensions[row_num2].height = 20
        row_num2 += 1

    ws2.merge_cells(start_row=row_num2, start_column=1, end_row=row_num2, end_column=2)
    ws2.cell(row=row_num2, column=1, value='ИТОГО ПО БРЕНДАМ').alignment = Alignment(horizontal='right', vertical='center')

    for b_idx in range(3, len(matrix_headers) + 1):
        b_letter = get_column_letter(b_idx)
        tot_c = ws2.cell(row=row_num2, column=b_idx, value=f'=SUM({b_letter}{start_row2+1}:{b_letter}{row_num2-1})')
        tot_c.alignment = align_right
        tot_c.number_format = '#,##0'

    for col_idx in range(1, len(matrix_headers) + 1):
        c = ws2.cell(row=row_num2, column=col_idx)
        c.font = total_font
        c.fill = total_fill
        c.border = total_border
    ws2.row_dimensions[row_num2].height = 24

    ws2.auto_filter.ref = f'A{start_row2}:{get_column_letter(len(matrix_headers))}{row_num2-1}'
    ws2.freeze_panes = 'C5'

    ws2.column_dimensions['A'].width = 8
    ws2.column_dimensions['B'].width = 42
    for b_idx in range(3, len(matrix_headers)):
        col_letter = get_column_letter(b_idx)
        brand_name = sorted_brands[b_idx - 3]
        ws2.column_dimensions[col_letter].width = max(len(brand_name) + 3, 11)
    ws2.column_dimensions[get_column_letter(len(matrix_headers))].width = 16

    # Target file paths
    target_filename = "Отчет_Сделки_Сентябрь_2026.xlsx"
    out_file = os.path.join(base_dir, target_filename)
    wb.save(out_file)
    print(f"Report saved to: {out_file}")

    # Also save a copy with ASCII name just in case the OS/UI has trouble with Cyrillic
    ascii_filename = "September_2026_Deals_Report.xlsx"
    ascii_out = os.path.join(base_dir, ascii_filename)
    wb.save(ascii_out)
    print(f"ASCII copy saved to: {ascii_out}")

if __name__ == '__main__':
    build_report()
