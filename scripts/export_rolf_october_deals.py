import os
import sys
import re
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from scripts.etl.normalization import get_exact_val, normalize_brand, parse_custom_date
from scripts.etl.tabular import read_tabular_file

def export_rolf_deals():
    latest_deal_path = os.path.join(ROOT, 'raw_data', 'DEAL_20261008_77cb3c12_6ac730e8d0d3a.xls')
    datasets = read_tabular_file(latest_deal_path)
    rows = datasets[0][1] if datasets else []

    deals_by_id = {}
    for r in rows:
        comp = str(get_exact_val(r, 'КОМПАНИЯНАЗВАНИЕКОМПАНИИ', 'КОМПАНИЯ') or '').strip()
        if 'рольф' not in comp.lower():
            continue
        did = str(get_exact_val(r, 'ID', 'IDСДЕЛКИ') or '').strip()
        if not did:
            continue
        if did not in deals_by_id:
            deals_by_id[did] = []
        deals_by_id[did].append(r)

    aux_keywords = ('ДОП.', 'АВАНС', 'БРОНЬ', 'СТРАХОВ', 'КАСКО', 'ОСАГО', 'СЕРТИФИКАТ')

    deals_list = []
    for did, r_list in sorted(deals_by_id.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 0):
        car_row = None
        prepay_row = None
        for r in r_list:
            tovar = str(get_exact_val(r, 'ТОВАР') or '').strip()
            t_up = tovar.upper()
            if 'ВНЕСЕНИЕ АВАНСА' in t_up:
                prepay_row = r
            elif not any(kw in t_up for kw in aux_keywords) and not car_row:
                car_row = r
        if not car_row:
            car_row = r_list[0]

        tovar = str(get_exact_val(car_row, 'ТОВАР') or '').strip()
        vin = str(get_exact_val(car_row, 'VIN', 'VINНОМЕР') or '').strip()
        if not vin:
            words = tovar.split()
            if words and len(words[-1]) >= 11 and re.search(r'[A-Z0-9]', words[-1]):
                vin = words[-1]

        brand = normalize_brand(tovar, month='2026-10')
        if not brand:
            # Fallback brand detection
            t_low = tovar.lower()
            if 'geely' in t_low or 'джили' in t_low: brand = 'Geely'
            elif 'belgee' in t_low or 'белджи' in t_low: brand = 'Belgee'
            elif 'jaecoo' in t_low or 'джейку' in t_low or 'jeland' in t_low: brand = 'Jaecoo'
            elif 'omoda' in t_low or 'омода' in t_low: brand = 'Omoda'
            elif 'jetour' in t_low or 'джетур' in t_low: brand = 'Jetour'
            elif 'gac' in t_low or 'гак' in t_low: brand = 'GAC'
            else: brand = 'Прочие'

        model = tovar
        for b_word in [brand, 'Geely', 'GEELY', 'Belgee', 'BELGEE', 'JAECOO', 'Jaecoo', 'JELAND', 'Jeland', 'OMODA', 'Omoda', 'JETOUR', 'Jetour', 'GAC']:
            model = re.sub(re.escape(b_word), '', model, flags=re.IGNORECASE)
        if vin:
            model = model.replace(vin, '')
        model = model.strip(' ,-\t\r\n')
        if not model:
            model = brand

        comp = str(get_exact_val(car_row, 'КОМПАНИЯНАЗВАНИЕКОМПАНИИ', 'КОМПАНИЯ') or '').strip()
        stage = str(get_exact_val(car_row, 'СТАДИЯСДЕЛКИ') or '').strip()
        b2c = str(get_exact_val(car_row, 'ТИПСДЕЛКИB2C') or '').strip()
        mgr = str(get_exact_val(car_row, 'МЕНЕДЖЕРСДЕЛКИ') or '').strip()
        close_d = str(get_exact_val(car_row, 'ПРЕДПОЛАГАЕМАЯДАТАЗАКРЫТИЯ') or '').strip()
        prepay_d = str(get_exact_val(car_row, 'ДАТАВНЕСЕНИЯПРЕДОПЛАТЫРОЗНИЦА', 'ДАТАПОЛУЧЕНИЯАВАНСА') or (get_exact_val(prepay_row, 'ДАТАПОЛУЧЕНИЯАВАНСА') if prepay_row else '') or '').strip()

        raw_p = str(get_exact_val(car_row, 'ФИНАЛЬНАЯЦЕНАB2C', 'ЦЕНА') or 0).replace(' ', '').replace('\xa0', '').replace(',', '.')
        try: price = float(raw_p)
        except: price = 0.0

        raw_c = str(get_exact_val(car_row, 'КОМИССИЯСДЕЛКИРУБ') or 0).replace(' ', '').replace('\xa0', '').replace(',', '.')
        try: comm = float(raw_c)
        except: comm = 0.0

        is_sold = 'ЗАКРЫТО И РЕАЛИЗОВАН' in stage.upper()
        status = 'Продажа' if is_sold else 'В работе'

        deals_list.append({
            'id': did,
            'company': comp,
            'brand': brand,
            'model': model,
            'vin': vin or '—',
            'manager': mgr,
            'b2c_type': b2c,
            'stage': stage,
            'status': status,
            'close_date': close_d,
            'prepay_date': prepay_d or '—',
            'price': price,
            'comm': comm
        })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сделки РОЛЬФ Октябрь 2026"

    # Turn on grid lines
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=16, bold=True, color="1E293B")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="64748B")
    
    kpi_title_font = Font(name=font_family, size=9, bold=True, color="64748B")
    kpi_value_font = Font(name=font_family, size=14, bold=True, color="0F172A")
    kpi_sub_font = Font(name=font_family, size=8, color="94A3B8")

    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=9, color="1E293B")
    bold_data_font = Font(name=font_family, size=9, bold=True, color="1E293B")
    total_font = Font(name=font_family, size=10, bold=True, color="0F172A")

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    sold_badge_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    sold_badge_font = Font(name=font_family, size=9, bold=True, color="166534")
    work_badge_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    work_badge_font = Font(name=font_family, size=9, bold=True, color="92400E")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    total_border = Border(
        top=Side(style='thin', color='94A3B8'),
        bottom=Side(style='double', color='0F172A')
    )

    # 1. Sheet Title
    ws['B2'] = "Реестр сделок АО «РОЛЬФ» — Октябрь 2026"
    ws['B2'].font = title_font
    ws['B3'] = f"Выгружено из CRM Битрикс24 на {datetime.date.today().strftime('%d.%m.%Y')} | Всего сделок в выборке: {len(deals_list)}"
    ws['B3'].font = subtitle_font

    # 2. KPI Cards Block
    total_deals = len(deals_list)
    sold_deals = sum(1 for d in deals_list if d['status'] == 'Продажа')
    work_deals = total_deals - sold_deals
    total_price = sum(d['price'] for d in deals_list)
    total_comm = sum(d['comm'] for d in deals_list)

    kpis = [
        ("B5", "C6", "ВСЕГО СДЕЛОК", f"{total_deals} шт.", "Все стадии за октябрь"),
        ("D5", "E6", "РЕАЛИЗОВАНО (ПРОДАЖИ)", f"{sold_deals} шт.", f"Доля: {round(sold_deals/total_deals*100, 1)}%"),
        ("F5", "G6", "В РАБОТЕ / АВАНСЫ", f"{work_deals} шт.", f"Доля: {round(work_deals/total_deals*100, 1)}%"),
        ("H5", "I6", "ОБЩИЙ ОБОРОТ АВТО", f"{int(total_price):,} ₽".replace(',', ' '), "По финальной цене авто"),
        ("J5", "K6", "КОМИССИЯ B2C", f"{int(total_comm):,} ₽".replace(',', ' '), "Суммарный доход сети")
    ]

    for top_l, btm_r, k_title, k_val, k_sub in kpis:
        c1, r1 = top_l[0], int(top_l[1:])
        c2, r2 = btm_r[0], int(btm_r[1:])
        ws.merge_cells(f"{c1}{r1}:{c2}{r1}")
        ws.merge_cells(f"{c1}{r2}:{c2}{r2}")
        
        ws[f"{c1}{r1}"] = k_title
        ws[f"{c1}{r1}"].font = kpi_title_font
        ws[f"{c1}{r1}"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws[f"{c1}{r2}"] = k_val
        ws[f"{c1}{r2}"].font = kpi_value_font
        ws[f"{c1}{r2}"].alignment = Alignment(horizontal="center", vertical="center")

        for r_idx in range(r1, r2 + 1):
            for c_col in [c1, c2]:
                cell = ws[f"{c_col}{r_idx}"]
                cell.fill = kpi_fill
                cell.border = thin_border

    # 3. Table Headers
    headers = [
        ("№", "center"),
        ("ID сделки", "center"),
        ("Статус", "center"),
        ("Марка", "left"),
        ("Модель автомобиля", "left"),
        ("VIN номер", "center"),
        ("Менеджер сделки", "left"),
        ("Тип сделки (B2C)", "center"),
        ("Стадия сделки в CRM", "left"),
        ("Дата закрытия", "center"),
        ("Дата аванса", "center"),
        ("Цена автомобиля (руб.)", "right"),
        ("Комиссия (руб.)", "right"),
        ("Юридическое лицо / Дилерский центр", "left")
    ]

    start_row = 8
    start_col = 2  # Column B

    for idx, (h_title, h_align) in enumerate(headers):
        col_letter = get_column_letter(start_col + idx)
        cell = ws[f"{col_letter}{start_row}"]
        cell.value = h_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 28

    # 4. Table Data Rows
    current_row = start_row + 1
    for i, d in enumerate(deals_list, 1):
        row_vals = [
            (i, "center", False, None),
            (int(d['id']) if d['id'].isdigit() else d['id'], "center", True, None),
            (d['status'], "center", True, sold_badge_fill if d['status'] == 'Продажа' else work_badge_fill),
            (d['brand'], "left", True, None),
            (d['model'], "left", False, None),
            (d['vin'], "center", False, None),
            (d['manager'], "left", False, None),
            (d['b2c_type'], "center", False, None),
            (d['stage'], "left", False, None),
            (d['close_date'], "center", False, None),
            (d['prepay_date'], "center", False, None),
            (d['price'], "right", False, None),
            (d['comm'], "right", False, None),
            (d['company'], "left", False, None)
        ]

        is_even = (i % 2 == 0)
        default_fill = alt_row_fill if is_even else None

        for col_idx, (val, align, is_bold, custom_fill) in enumerate(row_vals):
            col_letter = get_column_letter(start_col + col_idx)
            cell = ws[f"{col_letter}{current_row}"]
            cell.value = val

            # Styling
            if custom_fill:
                cell.fill = custom_fill
                cell.font = sold_badge_font if d['status'] == 'Продажа' else work_badge_font
            else:
                if default_fill:
                    cell.fill = default_fill
                cell.font = bold_data_font if is_bold else data_font

            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = thin_border

            # Number formatting
            if col_idx in (11, 12):  # Price and Commission
                cell.number_format = '#,##0.00 ₽'
            elif col_idx == 0:
                cell.number_format = '0'

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # 5. Summary / Total Row
    total_row = current_row
    ws.row_dimensions[total_row].height = 24

    ws[f"B{total_row}"] = ""
    ws[f"C{total_row}"] = "ИТОГО"
    ws[f"C{total_row}"].font = total_font
    ws[f"C{total_row}"].alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(len(headers)):
        col_letter = get_column_letter(start_col + col_idx)
        cell = ws[f"{col_letter}{total_row}"]
        cell.fill = total_fill
        cell.border = total_border
        cell.font = total_font

    # Counts and Sum formulas
    price_col = get_column_letter(start_col + 11)
    comm_col = get_column_letter(start_col + 12)
    
    ws[f"{price_col}{total_row}"] = f"=SUM({price_col}{start_row + 1}:{price_col}{total_row - 1})"
    ws[f"{price_col}{total_row}"].number_format = '#,##0.00 ₽'
    ws[f"{price_col}{total_row}"].alignment = Alignment(horizontal="right", vertical="center")

    ws[f"{comm_col}{total_row}"] = f"=SUM({comm_col}{start_row + 1}:{comm_col}{total_row - 1})"
    ws[f"{comm_col}{total_row}"].number_format = '#,##0.00 ₽'
    ws[f"{comm_col}{total_row}"].alignment = Alignment(horizontal="right", vertical="center")

    # 6. Auto-fit column widths
    padding = 3
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column < start_col or col[0].column >= start_col + len(headers):
            continue
        max_len = 0
        for cell in col:
            if cell.row < start_row:
                continue
            val_str = str(cell.value or '')
            if cell.number_format and ('₽' in cell.number_format or '#' in cell.number_format) and isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.2f} ₽"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + padding, 12)

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 6   # №
    ws.column_dimensions['C'].width = 13  # ID сделки
    ws.column_dimensions['D'].width = 14  # Статус
    ws.column_dimensions['E'].width = 16  # Марка
    ws.column_dimensions['F'].width = 24  # Модель
    ws.column_dimensions['G'].width = 22  # VIN
    ws.column_dimensions['H'].width = 22  # Менеджер
    ws.column_dimensions['I'].width = 15  # Тип сделки
    ws.column_dimensions['J'].width = 24  # Стадия
    ws.column_dimensions['K'].width = 15  # Дата закрытия
    ws.column_dimensions['L'].width = 15  # Дата аванса
    ws.column_dimensions['M'].width = 18  # Цена
    ws.column_dimensions['N'].width = 16  # Комиссия
    ws.column_dimensions['O'].width = 32  # Компания

    # Enable auto filter
    ws.auto_filter.ref = f"B{start_row}:O{total_row - 1}"

    # Target path in Downloads
    out_dir = r"C:\Users\pc\Downloads"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Сделки_РОЛЬФ_Октябрь_2026.xlsx")
    wb.save(out_file)
    print(f"[*] Файл успешно сохранен: {out_file}")
    return out_file

if __name__ == '__main__':
    export_rolf_deals()
