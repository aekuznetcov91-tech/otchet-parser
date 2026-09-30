import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('raw_data/OEM СберАвто финал (19).xlsx', data_only=True)

TYPO_FIX = {
    'санк-петербург': 'Санкт-Петербург',
    'санкт-петебург': 'Санкт-Петербург',
    'новоросийск': 'Новороссийск',
    'набережный челны': 'Набережные Челны',
    'мин. воды': 'Минеральные Воды',
    'минеральные воды': 'Минеральные Воды',
    'уфа': 'Уфа',
    'ростов-на-дону': 'Ростов-на-Дону',
    'нижний новгород': 'Нижний Новгород',
    'дзержинск, нижегородская область': 'Дзержинск (Нижегородская обл.)',
    'новомосковск (тульская область)': 'Новомосковск (Тульская обл.)',
    'оренбургская обл., г. бузулук': 'Бузулук (Оренбургская обл.)',
    'самарская обл., г. новокуйбышевск': 'Новокуйбышевск (Самарская обл.)',
    'самарская обл., г. сызрань': 'Сызрань (Самарская обл.)',
    'челябинск область': 'Челябинская область',
    'ставропольский край': 'Ставропольский край',
}

all_cities = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    city_col = None
    header_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column + 1, 40)):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if val.lower() == 'город':
                city_col = c
                header_row = r
                break
        if city_col is not None:
            break
            
    if city_col is None:
        continue
        
    for r in range(header_row + 1, ws.max_row + 1):
        c_val = ws.cell(row=r, column=city_col).value
        if c_val is not None and str(c_val).strip():
            c_str = str(c_val).strip()
            if c_str.lower().startswith(('итого', 'примечание', 'всего', 'город')):
                continue
            c_low = c_str.lower()
            
            if c_low in TYPO_FIX:
                canon = TYPO_FIX[c_low]
            else:
                canon = c_str.capitalize() if c_str.isupper() and len(c_str) > 3 else c_str
                
            key = canon.lower()
            if key not in all_cities:
                all_cities[key] = {
                    'name': canon,
                    'sheets': set(),
                    'raw_variants': set()
                }
            all_cities[key]['sheets'].add(sname)
            all_cities[key]['raw_variants'].add(c_str)

sorted_all = sorted(list(all_cities.values()), key=lambda x: x['name'].lower())

with open('all_unique_cities_final.txt', 'w', encoding='utf-8') as f:
    f.write(f"ИТОГО УНИКАЛЬНЫХ ГОРОДОВ: {len(sorted_all)}\n\n")
    for i, item in enumerate(sorted_all, 1):
        c_name = item['name']
        sheets = sorted(list(item['sheets']))
        sheets_str = ', '.join(sheets)
        f.write(f"{i}. {c_name} — {len(sheets)} бр.: {sheets_str}\n")

print("Done! Total cities:", len(sorted_all))
