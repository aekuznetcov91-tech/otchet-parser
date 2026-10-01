import os
import re
import json
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Paths
html_path = '/Users/vaceslavgamaunov/Downloads/Новые авто.html'
data_json_path = 'site/data.json'
output_desktop = '/Users/vaceslavgamaunov/Desktop/Ранжирование_автомобилей_витрины_СберАвто_Carsale.xlsx'
output_downloads = '/Users/vaceslavgamaunov/Downloads/Ранжирование_автомобилей_витрины_СберАвто_Carsale.xlsx'

with open(html_path, 'r', encoding='utf-8') as f:
    html = f.read()

pattern = re.compile(r'<a[^>]+data-testid=[\"\']CarsPage\.List\.MiniAutoCard\.([^\"\']+)[\"\'][^>]*href=[\"\']([^\"\']+)[\"\'][^>]*>(.*?)</a>', re.DOTALL)
matches = pattern.findall(html)

all_cards = []
for idx, (testid, href, inner) in enumerate(matches, 1):
    clean = re.sub(r'<[^>]+>', '\n', inner)
    lines = [l.strip().replace('&nbsp;', ' ') for l in clean.split('\n') if l.strip()]
    m_name = testid.rsplit('.', 1)[0]
    price = None
    price_str = ''
    discount = ''
    dealers = ''
    specs = ''
    for l in lines:
        if 'Скидка' in l: discount = l
        elif '₽' in l and not price:
            p_clean = re.sub(r'[^\d]', '', l)
            if p_clean and int(p_clean) > 500000:
                price = int(p_clean)
                price_str = l
        elif 'авто от' in l: dealers = l
        elif any(k in l for k in ['л,', 'кВт,', 'Механика', 'Автомат', 'Робот', 'Вариатор']): specs = l
            
    all_cards.append({
        'pos': idx, 'name': m_name, 'testid': testid, 'href': href,
        'price': price, 'price_str': price_str, 'discount': discount,
        'dealers': dealers, 'specs': specs
    })

def get_card_family(name):
    n = name.upper()
    if 'GRANTA' in n: return 'LADA Granta'
    if 'VESTA' in n: return 'LADA Vesta'
    if 'NIVA TRAVEL' in n: return 'LADA Niva Travel'
    if 'NIVA LEGEND' in n: return 'LADA Niva Legend'
    if 'NIVA' in n: return 'LADA Niva (Legend / Travel)'
    if 'LARGUS' in n: return 'LADA Largus'
    if 'ISKRA' in n: return 'LADA Iskra'
    if 'AURA' in n: return 'LADA Aura'
    if 'DASHING' in n: return 'Jetour Dashing'
    if 'X70' in n: return 'Jetour X70 PLUS'
    if 'X90' in n: return 'Jetour X90 PLUS'
    if 'T1' in n or 'Т1' in n: return 'Jetour T1'
    if 'T2' in n or 'Т2' in n: return 'Jetour T2'
    if 'SOUEAST S06' in n or 'S06' in n: return 'Soueast S06'
    if 'SOUEAST S07' in n or 'S07' in n: return 'Soueast S07'
    if 'GS8' in n: return 'GAC GS8'
    if 'GS4' in n: return 'GAC GS4'
    if 'GS3' in n: return 'GAC GS3'
    if 'EMPOW' in n: return 'GAC Empow'
    if 'AION' in n: return 'GAC Aion V'
    if 'M8' in n: return 'GAC M8'
    if 'S7' in n: return 'GAC S7'
    if 'S9' in n: return 'GAC S9'
    if 'JOLION' in n: return 'Haval Jolion'
    if 'M6' in n: return 'Haval M6'
    if 'H3' in n: return 'Haval H3'
    if 'H5' in n: return 'Haval H5'
    if 'H7' in n: return 'Haval H7'
    if 'H9' in n: return 'Haval H9'
    if 'F7X' in n: return 'Haval F7x'
    if 'F7' in n: return 'Haval F7'
    if 'DARGO' in n: return 'Haval Dargo'
    if 'SOLARIS HC' in n: return 'Solaris HC'
    if 'SOLARIS KRX' in n: return 'Solaris KRX'
    if 'SOLARIS KRS' in n: return 'Solaris KRS'
    if 'MONJARO' in n: return 'Geely Monjaro'
    if 'COOLRAY' in n: return 'Geely Coolray'
    if 'CITYRAY' in n: return 'Geely Cityray'
    if 'ATLAS' in n: return 'Geely Atlas'
    if 'PREFACE' in n: return 'Geely Preface'
    if 'OKAVANGO' in n: return 'Geely Okavango'
    if 'EX5' in n: return 'Geely EX5'
    if 'BELGEE X50' in n or 'X50+' in n: return 'Belgee X50+'
    if 'BELGEE X70' in n or 'Х70' in n: return 'Belgee X70'
    if 'BELGEE S50' in n or 'S50' in n: return 'Belgee S50'
    if 'TENET T4L' in n: return 'Tenet T4L'
    if 'TENET T4' in n: return 'Tenet T4'
    if 'TENET T7' in n: return 'Tenet T7'
    if 'TENET T8' in n: return 'Tenet T8'
    if 'TENET T9' in n: return 'Tenet T9'
    if 'TENET A8' in n: return 'Tenet A8'
    if 'JELAND J6' in n: return 'Jeland J6'
    if 'JELAND J7' in n: return 'Jeland J7'
    if 'JAECOO J6' in n: return 'Jaecoo J6'
    if 'JAECOO J7' in n: return 'Jaecoo J7'
    if 'JAECOO J8' in n: return 'Jaecoo J8'
    if 'OMODA C5' in n: return 'Omoda C5'
    if 'OMODA C7' in n: return 'Omoda C7'
    if 'UNI-S' in n: return 'Changan UNI-S'
    if 'UNI-V' in n: return 'Changan UNI-V'
    if 'UNI-K' in n: return 'Changan UNI-K'
    if 'UNI-T' in n: return 'Changan UNI-T'
    if 'CS35' in n: return 'Changan CS35PLUS / CS35MAX'
    if 'CS75' in n: return 'Changan CS75PLUS / CS75PRO'
    if 'CS95' in n: return 'Changan CS95'
    if 'ALSVIN' in n: return 'Changan Alsvin'
    if 'EADO' in n: return 'Changan Eado Plus'
    if 'HUNTER' in n: return 'Changan Hunter Plus'
    if 'DEEPAL' in n or 'G318' in n: return 'Changan Deepal G318'
    if 'МОСКВИЧ 3' in n: return 'Москвич 3'
    if 'МОСКВИЧ 8' in n: return 'Москвич 8'
    if 'М70' in n: return 'Москвич М70'
    if 'М90' in n: return 'Москвич М90'
    if 'MAZDA' in n: return 'Mazda CX-5'
    if 'TOYOTA' in n or 'CAMRY' in n: return 'Toyota Camry'
    if 'SANTA FE' in n or 'HYUNDAI' in n: return 'Hyundai Santa Fe'
    return name

for c in all_cards:
    c['family'] = get_card_family(c['name'])

showcase_families_set = sorted(list(set(c['family'] for c in all_cards)))

def match_deal(b_raw, m_raw):
    b = str(b_raw).strip().upper()
    m = str(m_raw).strip().upper()
    
    if 'JETOUR' in b:
        if 'DASH' in m: return 'Jetour Dashing'
        if any(x in m for x in ['T1', 'Т1', 'T-1', 'TI']): return 'Jetour T1'
        if any(x in m for x in ['T2', 'Т2', 'T-2']): return 'Jetour T2'
        if any(x in m for x in ['X70', 'Х70', '70+', '70 +']): return 'Jetour X70 PLUS'
        if any(x in m for x in ['X90', 'Х90']): return 'Jetour X90 PLUS'
        return 'Jetour Другое'
        
    if 'LADA' in b or 'ВАЗ' in b:
        if 'GRANT' in m or 'GANT' in m: return 'LADA Granta'
        if 'VEST' in m or 'GFL' in m: return 'LADA Vesta'
        if 'TRAVEL' in m or 'TREVEL' in m: return 'LADA Niva Travel'
        if 'LEGEND' in m: return 'LADA Niva Legend'
        if any(x in m for x in ['NIVA', 'НИВА', '4X4']): return 'LADA Niva Legend'
        if 'LARG' in m: return 'LADA Largus'
        if 'ISKR' in m: return 'LADA Iskra'
        if 'AURA' in m: return 'LADA Aura'
        return 'LADA Другое'

    if 'SOUEAST' in b or 'SOUEAS' in b:
        if '06' in m: return 'Soueast S06'
        if '07' in m or 'С07' in m: return 'Soueast S07'
        return 'Soueast Другое'

    if 'GAC' in b:
        if 'GS8' in m or 'GS 8' in m or 'CS8' in m: return 'GAC GS8'
        if 'GS4' in m or 'GS 4' in m or 'CS4' in m: return 'GAC GS4'
        if 'M8' in m or 'М8' in m: return 'GAC M8'
        if 'S7' in m: return 'GAC S7'
        if 'S9' in m or 'GS9' in m: return 'GAC S9'
        if 'GS3' in m: return 'GAC GS3'
        if 'EMPOW' in m: return 'GAC Empow'
        if 'AION' in m: return 'GAC Aion V'
        return 'GAC Другое'

    if 'HAVAL' in b or 'ХАВЕЙЛ' in b or 'ХАВАЛ' in b:
        if 'JOLION' in m or 'ДЖОЛИОН' in m: return 'Haval Jolion'
        if 'M6' in m or 'М6' in m: return 'Haval M6'
        if 'F7X' in m: return 'Haval F7x'
        if 'F7' in m: return 'Haval F7'
        if 'DARGO' in m or 'ДАРГО' in m: return 'Haval Dargo'
        if 'H3' in m: return 'Haval H3'
        if 'H5' in m: return 'Haval H5'
        if 'H7' in m: return 'Haval H7'
        if 'H9' in m: return 'Haval H9'
        return 'Haval Другое'

    if 'GEELY' in b or 'ДЖИЛИ' in b:
        if 'MONJARO' in m or 'МАНДЖАРО' in m: return 'Geely Monjaro'
        if 'CITYRAY' in m or 'СИТИРЕЙ' in m: return 'Geely Cityray'
        if 'COOLRAY' in m or 'КУЛРЕЙ' in m: return 'Geely Coolray'
        if 'ATLAS' in m or 'АТЛАС' in m: return 'Geely Atlas'
        if 'PREFACE' in m or 'ПРЕФЕЙС' in m: return 'Geely Preface'
        if 'OKAVANGO' in m or 'ОКАВАНГО' in m: return 'Geely Okavango'
        if 'EX5' in m: return 'Geely EX5'
        return 'Geely Другое'

    if 'BELGEE' in b or 'БЕЛДЖИ' in b:
        if '50' in m and 'S50' not in m: return 'Belgee X50+'
        if '70' in m: return 'Belgee X70'
        if 'S50' in m: return 'Belgee S50'
        return 'Belgee Другое'

    if 'SOLARIS' in b or 'СОЛЯРИС' in b:
        if 'HC' in m or 'НС' in m: return 'Solaris HC'
        if 'KRX' in m: return 'Solaris KRX'
        if 'KRS' in m: return 'Solaris KRS'
        return 'Solaris Другое'

    if 'CHANGAN' in b or 'ЧАНГАН' in b:
        if 'UNI-S' in m or 'UNI S' in m or 'CS55' in m: return 'Changan UNI-S'
        if 'UNI-V' in m or 'UNI V' in m: return 'Changan UNI-V'
        if 'UNI-K' in m or 'UNI K' in m: return 'Changan UNI-K'
        if 'UNI-T' in m: return 'Changan UNI-T'
        if '35' in m: return 'Changan CS35PLUS / CS35MAX'
        if '75' in m: return 'Changan CS75PLUS / CS75PRO'
        if '95' in m: return 'Changan CS95'
        if 'ALSVIN' in m or 'АЛСВИН' in m: return 'Changan Alsvin'
        if 'EADO' in m or 'ИДО' in m: return 'Changan Eado Plus'
        if 'HUNTER' in m or 'ХАНТЕР' in m: return 'Changan Hunter Plus'
        if 'DEEPAL' in m or 'G318' in m: return 'Changan Deepal G318'
        return 'Changan Другое'

    if 'TENET' in b:
        if 'T4L' in m or 'Т4L' in m or 'T4 L' in m: return 'Tenet T4L'
        if 'T4' in m or 'Т4' in m: return 'Tenet T4'
        if 'T7' in m or 'Т7' in m: return 'Tenet T7'
        if 'T8' in m or 'Т8' in m: return 'Tenet T8'
        if 'T9' in m or 'Т9' in m: return 'Tenet T9'
        if 'A8' in m or 'А8' in m: return 'Tenet A8'
        return 'Tenet Другое'

    if 'JELAND' in b:
        if 'J6' in m: return 'Jeland J6'
        if 'J7' in m: return 'Jeland J7'
        return 'Jeland Другое'
    if 'JAECOO' in b:
        if 'J6' in m: return 'Jaecoo J6'
        if 'J7' in m: return 'Jaecoo J7'
        if 'J8' in m: return 'Jaecoo J8'
        return 'Jaecoo Другое'

    if 'OMODA' in b:
        if 'C5' in m: return 'Omoda C5'
        if 'C7' in m: return 'Omoda C7'
        return 'Omoda Другое'

    if 'МОСКВИЧ' in b or 'MOSKVICH' in b:
        if '3' in m: return 'Москвич 3'
        if '8' in m: return 'Москвич 8'
        if '70' in m or 'М70' in m: return 'Москвич М70'
        if '90' in m or 'М90' in m: return 'Москвич М90'
        return 'Москвич Другое'

    if 'TOYOTA' in b or 'ТОЙОТА' in b: return 'Toyota Camry'
    if 'MAZDA' in b or 'МАЗДА' in b: return 'Mazda CX-5'
    if 'HYUNDAI' in b or 'ХЕНДЭ' in b or 'ХЮНДАЙ' in b: return 'Hyundai Santa Fe'

    return 'Прочее'

with open(data_json_path, 'r', encoding='utf-8') as f:
    d = json.load(f)
sys_db = d['sys_db']
retail_b2c = {'ФДЦ', 'ФДЦ+ГП', 'Online', 'Передача лида'}

fam_stats = defaultdict(lambda: {
    'jul_tot': 0, 'jul_ret': 0,
    'aug_tot': 0, 'aug_ret': 0,
    'sep_tot': 0, 'sep_ret': 0,
    'q3_tot': 0, 'q3_ret': 0,
    'rev': 0.0, 'prices': []
})

q3_months = {'2026-07', '2026-08', '2026-09'}

for r in sys_db:
    if r.get('SaleQty', 0) != 1: continue
    m = r.get('SaleMonth')
    if m not in q3_months: continue
    b = str(r.get('Brand') or '').strip()
    md = str(r.get('Model') or '').strip()
    p = r.get('Price', 0) or 0
    rev = r.get('Revenue', 0) or 0
    is_ret = r.get('B2C') in retail_b2c
    
    fam = match_deal(b, md)
    fam_stats[fam]['q3_tot'] += 1
    fam_stats[fam]['rev'] += rev
    if is_ret: fam_stats[fam]['q3_ret'] += 1
    if p > 0: fam_stats[fam]['prices'].append(p)
    
    if m == '2026-07':
        fam_stats[fam]['jul_tot'] += 1
        if is_ret: fam_stats[fam]['jul_ret'] += 1
    elif m == '2026-08':
        fam_stats[fam]['aug_tot'] += 1
        if is_ret: fam_stats[fam]['aug_ret'] += 1
    elif m == '2026-09':
        fam_stats[fam]['sep_tot'] += 1
        if is_ret: fam_stats[fam]['sep_ret'] += 1

ranked = []
for fam in showcase_families_set:
    s = fam_stats[fam]
    tot = s['q3_tot']
    ret = s['q3_ret']
    rev = s['rev']
    avg_price = sum(s['prices']) / len(s['prices']) if s['prices'] else 0
    showcase_prices = [c['price'] for c in all_cards if c['family'] == fam and c['price']]
    min_price = min(showcase_prices) if showcase_prices else avg_price
    
    # Tier classification for Q3 (3 months)
    if tot >= 100:
        tier = "Тир 1: Суперхиты витрины (100+ сделок)"
        tier_code = 1
    elif tot >= 30:
        tier = "Тир 2: Высокий спрос (30–99 сделок)"
        tier_code = 2
    elif tot >= 10:
        tier = "Тир 3: Умеренные продажи (10–29 сделок)"
        tier_code = 3
    elif tot >= 1:
        tier = "Тир 4: Единичные продажи (1–9 сделок)"
        tier_code = 4
    else:
        tier = "Тир 5: Нулевые продажи (0 сделок в Q3)"
        tier_code = 5

    ranked.append({
        'family': fam,
        'tier': tier,
        'tier_code': tier_code,
        'q3_tot': tot,
        'q3_ret': ret,
        'jul_tot': s['jul_tot'],
        'jul_ret': s['jul_ret'],
        'aug_tot': s['aug_tot'],
        'aug_ret': s['aug_ret'],
        'sep_tot': s['sep_tot'],
        'sep_ret': s['sep_ret'],
        'rev': rev,
        'avg_price': avg_price,
        'min_price': min_price
    })

ranked.sort(key=lambda x: (x['tier_code'], -x['q3_tot'], -x['q3_ret'], -x['rev']))

# Also assign target positions to all 148 cards
card_target_pos = []
for idx, r in enumerate(ranked, 1):
    fam = r['family']
    matching_cards = [c for c in all_cards if c['family'] == fam]
    # sort matching cards by price asc
    matching_cards.sort(key=lambda x: x['price'] or 999999999)
    for c in matching_cards:
        c['target_family_rank'] = idx
        c['tier'] = r['tier']
        card_target_pos.append(c)

# Assign sequential target_pos 1..148
for i, c in enumerate(card_target_pos, 1):
    c['target_pos'] = i

print(f"[*] Сформирован актуальный рейтинг Q3 из {len(ranked)} модельных линеек для {len(card_target_pos)} карточек.")

# Save Excel
wb = openpyxl.Workbook()
# Sheet 1: Ranking
ws1 = wb.active
ws1.title = "Ранжирование Q3 (73 модели)"
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    "Целевая позиция", "Тир спроса", "Модель на витрине", 
    "Сделки Q3 (Всего)", "В т.ч. Розница B2C", "Доля розницы, %",
    "Июль (Все / Розн)", "Август (Все / Розн)", "Сентябрь (Все / Розн)",
    "Стартовая цена витрины, ₽", "Средняя цена сделки, ₽", "Выручка комиссии СберАвто (КВ), ₽"
]
ws1.append(headers1)

for col_idx in range(1, len(headers1) + 1):
    cell = ws1.cell(row=1, column=col_idx)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for idx, r in enumerate(ranked, 1):
    ret_share = (r['q3_ret'] / r['q3_tot'] * 100) if r['q3_tot'] else 0
    row_data = [
        idx, r['tier'], r['family'],
        r['q3_tot'], r['q3_ret'], f"{ret_share:.1f}%",
        f"{r['jul_tot']} / {r['jul_ret']}",
        f"{r['aug_tot']} / {r['aug_ret']}",
        f"{r['sep_tot']} / {r['sep_ret']}",
        r['min_price'], r['avg_price'], r['rev']
    ]
    ws1.append(row_data)

# Sheet 2: All 148 cards
ws2 = wb.create_sheet(title="Все 148 карточек витрины")
ws2.views.sheetView[0].showGridLines = True
headers2 = [
    "Целевая позиция (Top->Bottom)", "Текущая позиция на витрине", "Смещение позиции",
    "Тир", "Семейство модели", "Название карточки", "Цена на витрине, ₽", 
    "Скидка", "Характеристики / ДЦ", "Ссылка на витрину Carsale"
]
ws2.append(headers2)

for col_idx in range(1, len(headers2) + 1):
    cell = ws2.cell(row=1, column=col_idx)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for c in card_target_pos:
    diff = c['pos'] - c['target_pos']
    diff_str = f"↑ +{diff}" if diff > 0 else (f"↓ {diff}" if diff < 0 else "=")
    row_c = [
        c['target_pos'], c['pos'], diff_str,
        c['tier'], c['family'], c['name'], c['price'],
        c['discount'], f"{c['specs']} | {c['dealers']}",
        f"https://appsb.sberauto.com{c['href']}"
    ]
    ws2.append(row_c)

# Auto-adjust column widths
for ws in [ws1, ws2]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

wb.save(output_desktop)
wb.save(output_downloads)
print(f"[+] Успешно сохранено в {output_desktop} и {output_downloads}")
