import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('raw_data/OEM СберАвто финал (19).xlsx', data_only=True)

# Normalization mapping for obvious typos
TYPO_FIX = {
    'санк-петербург': 'Санкт-Петербург',
    'санкт-петебург': 'Санкт-Петербург',
    'новоросийск': 'Новороссийск',
    'набережный челны': 'Набережные Челны',
    'мин. воды': 'Минеральные Воды',
    'уфа': 'Уфа',
    'ростов-на-дону': 'Ростов-на-Дону',
    'нижний новгород': 'Нижний Новгород',
}

results = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    city_col = None
    header_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column + 1, 40)):
            val = str(ws.cell(row=r, column=c).value or '').strip()
            if val.lower() == 'город':
                city_col = c
                header_row = r
                break
        if city_col is not None:
            break
            
    if city_col is None:
        results[sname] = {
            'has_table': False,
            'reason': 'Вкладка не содержит таблицы дилерских центров (столбец "Город" отсутствует)'
        }
        continue
        
    cities_map = {}
    for r in range(header_row + 1, ws.max_row + 1):
        c_val = ws.cell(row=r, column=city_col).value
        if c_val is not None and str(c_val).strip():
            c_str = str(c_val).strip()
            if c_str.lower().startswith(('итого', 'примечание', 'всего', 'город')):
                continue
            c_low = c_str.lower()
            
            # Canonical name
            if c_low in TYPO_FIX:
                canon = TYPO_FIX[c_low]
            else:
                canon = c_str.capitalize() if c_str.isupper() and len(c_str) > 3 else c_str
                
            canon_key = canon.lower()
            if canon_key not in cities_map:
                cities_map[canon_key] = {
                    'name': canon,
                    'raw_variants': set()
                }
            cities_map[canon_key]['raw_variants'].add(c_str)
            
    sorted_cities = sorted(list(cities_map.values()), key=lambda x: x['name'].lower())
    results[sname] = {
        'has_table': True,
        'count': len(sorted_cities),
        'cities': sorted_cities
    }

with open('all_brands_cities_summary.txt', 'w', encoding='utf-8') as f:
    for sname, data in results.items():
        f.write(f"### Вкладка `{sname}`\n")
        if not data['has_table']:
            f.write(f"*{data['reason']}*\n\n")
            continue
        if data['count'] == 0:
            f.write("*Таблица дилеров пуста (нет записей)*\n\n")
            continue
        f.write(f"**Уникальных городов: {data['count']}**\n\n")
        for i, item in enumerate(data['cities'], 1):
            c_name = item['name']
            vars_set = item['raw_variants']
            # note if typo was fixed
            extra = ""
            if len(vars_set) > 1 or (len(vars_set) == 1 and list(vars_set)[0] != c_name):
                extra = f" *(в файле: {', '.join(sorted(vars_set))})*"
            f.write(f"{i}. {c_name}{extra}\n")
        f.write("\n")

print("Generated all_brands_cities_summary.txt successfully!")
