import os
import sys
import re
import json
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Set stdout encoding
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from scripts.etl.tabular import read_tabular_file

def resolve_clean_city(raw_text, partner_name=''):
    if not raw_text or str(raw_text).lower() in ('nan', 'none', 'null', 'не указан', ''):
        raw_text = ''
    t = str(raw_text).strip()
    p = str(partner_name).strip()
    combined = (t + ' ' + p).upper()

    # Direct city keywords
    if any(k in combined for k in ['МОСКВ', 'ЗЕЛЕНОГРАД', 'ХИМКИ', 'МКАД', 'БАЛАШИХ', 'КРАСНОГОРСК', 'ЛЮБЕРЦ', 'ОДИНЦОВО', 'МЫТИЩ', 'КОТЕЛЬНИК', 'РЕУТОВ', '143420', 'ДЗЕРЖИНСКИЙ', 'ВИДНОЕ', 'ДОЛГОПРУДН']):
        return 'Москва'
    if any(k in combined for k in ['САНКТ-ПЕТЕРБУРГ', 'ПЕТЕРБУРГ', 'СПБ', 'ЛАХТА', 'ВЫБОРГ', 'ГАТЧИН']):
        return 'Санкт-Петербург'
    if any(k in combined for k in ['ЕКАТЕРИНБУРГ', 'СВЕРДЛОВСК', 'ЕКБ', 'ТАГИЛ']):
        return 'Екатеринбург'
    if any(k in combined for k in ['НИЖНИЙ НОВГОРОД', 'Н.НОВГОРОД', 'НИЖЕГОРОД']) or combined == 'НН':
        return 'Нижний Новгород'
    if any(k in combined for k in ['КРАСНОДАР', 'АНАПА', 'НОВОРОССИЙСК', 'АРМАВИР']):
        return 'Краснодар'
    if any(k in combined for k in ['КАЗАН']):
        return 'Казань'
    if any(k in combined for k in ['САМАР', 'ТОЛЬЯТТИ', 'СЫЗРАН']):
        return 'Самара'
    if any(k in combined for k in ['ВОРОНЕЖ', 'БОРИСОГЛЕБСК']):
        return 'Воронеж'
    if any(k in combined for k in ['РОСТОВ-НА-ДОНУ', 'РОСТОВ НА ДОНУ', 'РОСТОВ', 'ТАГАНРОГ', 'ШАХТЫ', 'БАТАЙСК']):
        return 'Ростов-на-Дону'
    if any(k in combined for k in ['ВОЛГОГРАД', 'ВОЛЖСК']):
        return 'Волгоград'
    if any(k in combined for k in ['НОВОСИБИРСК', 'БЕРДСК']) or combined == 'НСК':
        return 'Новосибирск'
    if any(k in combined for k in ['УФА', 'СТЕРЛИТАМАК', 'САЛАВАТ']):
        return 'Уфа'
    if any(k in combined for k in ['ПЕРМ', 'БЕРЕЗНИК']):
        return 'Пермь'
    if any(k in combined for k in ['ТЮМЕН', 'ТОБОЛЬСК']):
        return 'Тюмень'
    if any(k in combined for k in ['ЧЕЛЯБИНСК', 'МАГНИТОГОРСК', 'МИАСС']):
        return 'Челябинск'
    if any(k in combined for k in ['НАБЕРЕЖНЫЕ ЧЕЛНЫ', 'ЧЕЛНЫ', 'АЛЬМЕТЬЕВ', 'НИЖНЕКАМ']):
        return 'Набережные Челны'
    if any(k in combined for k in ['СТАВРОПОЛЬ', 'КУЛАКОВА']):
        return 'Ставрополь'
    if any(k in combined for k in ['МИНЕРАЛЬНЫЕ ВОДЫ', 'МИН ВОДЫ', 'МИНВОДЫ', 'ПЯТИГОРСК', 'КИСЛОВОДСК', 'ЕССЕНТУК']):
        return 'Минеральные Воды'
    if any(k in combined for k in ['ОРЕНБУРГ', 'ОРСК']):
        return 'Оренбург'
    if any(k in combined for k in ['ВЛАДИМИР', 'КОВРОВ', 'МУРОМ']):
        return 'Владимир'
    if any(k in combined for k in ['ТВЕР', 'РЖЕВ']):
        return 'Тверь'
    if any(k in combined for k in ['САРАТОВ', 'ЭНГЕЛЬС', 'БАЛАКОВО']):
        return 'Саратов'
    if any(k in combined for k in ['ЯРОСЛАВ', 'РЫБИНСК']):
        return 'Ярославль'
    if any(k in combined for k in ['РЯЗАН']):
        return 'Рязань'
    if any(k in combined for k in ['ТУЛА', 'НОВОМОСКОВСК']):
        return 'Тула'
    if any(k in combined for k in ['ИЖЕВСК', 'САРАПУЛ']):
        return 'Ижевск'
    if any(k in combined for k in ['КИРОВ']):
        return 'Киров'
    if any(k in combined for k in ['ЛИПЕЦК', 'ЕЛЕЦ']):
        return 'Липецк'
    if any(k in combined for k in ['БЕЛГОРОД', 'СТАРЫЙ ОСКОЛ', 'ГУБКИН']):
        return 'Белгород'
    if any(k in combined for k in ['КАЛУГА', 'ОБНИНСК']):
        return 'Калуга'
    if any(k in combined for k in ['КУРСК', 'ЖЕЛЕЗНОГОРСК']):
        return 'Курск'
    if any(k in combined for k in ['БРЯНСК']):
        return 'Брянск'
    if any(k in combined for k in ['ИВАНОВО', 'КИНЕШМА']):
        return 'Иваново'
    if any(k in combined for k in ['УЛЬЯНОВСК', 'ДИМИТРОВГРАД']):
        return 'Ульяновск'
    if any(k in combined for k in ['ЧЕБОКСАР', 'НОВОЧЕБОКСАРСК']):
        return 'Чебоксары'
    if any(k in combined for k in ['ПЕТРОЗАВОДСК']):
        return 'Петрозаводск'
    if any(k in combined for k in ['СУРГУТ', 'НИЖНЕВАРТОВСК', 'НЕФТЕЮГАНСК', 'ХМАО']):
        return 'Сургут'
    if any(k in combined for k in ['СОЧИ']):
        return 'Сочи'
    if any(k in combined for k in ['ВЕЛИКИЙ НОВГОРОД']):
        return 'Великий Новгород'
    if any(k in combined for k in ['ВЕЛИКИЕ ЛУКИ']):
        return 'Великие Луки'
    if any(k in combined for k in ['ПСКОВ']):
        return 'Псков'
    if any(k in combined for k in ['КАЛИНИНГРАД']):
        return 'Калининград'
    if any(k in combined for k in ['СИМФЕРОПОЛЬ', 'СЕВАСТОПОЛЬ', 'КРЫМ']):
        return 'Симферополь'
    if any(k in combined for k in ['АРХАНГЕЛЬСК', 'СЕВЕРОДВИНСК']):
        return 'Архангельск'
    if any(k in combined for k in ['МУРМАНСК', 'АПАТИТЫ']):
        return 'Мурманск'
    if any(k in combined for k in ['АСТРАХАН']):
        return 'Астрахань'
    if any(k in combined for k in ['ОМСК']):
        return 'Омск'
    if any(k in combined for k in ['КРАСНОЯРСК', 'НОРИЛЬСК', 'АЧИНСК']):
        return 'Красноярск'
    if any(k in combined for k in ['БАРНАУЛ', 'БИЙСК']):
        return 'Барнаул'
    if any(k in combined for k in ['ТОМСК']):
        return 'Томск'
    if any(k in combined for k in ['КЕМЕРОВО']):
        return 'Кемерово'
    if any(k in combined for k in ['НОВОКУЗНЕЦК', 'ПРОКОПЬЕВСК', 'КУЗБАСС']):
        return 'Новокузнецк'
    if any(k in combined for k in ['ИРКУТСК', 'БРАТСК', 'АНГАРСК']):
        return 'Иркутск'
    if any(k in combined for k in ['ХАБАРОВСК', 'КОМСОМОЛЬСК']):
        return 'Хабаровск'
    if any(k in combined for k in ['ВЛАДИВОСТОК', 'УССУРИЙСК', 'НАХОДКА']):
        return 'Владивосток'
    if any(k in combined for k in ['МАХАЧКАЛА', 'ДЕРБЕНТ']):
        return 'Махачкала'
    if any(k in combined for k in ['ТАМБОВ']):
        return 'Тамбов'
    if any(k in combined for k in ['ОРЕЛ']):
        return 'Орел'
    if any(k in combined for k in ['СМОЛЕНСК']):
        return 'Смоленск'
    if any(k in combined for k in ['КОСТРОМА']):
        return 'Кострома'
    if any(k in combined for k in ['ВОЛОГДА', 'ЧЕРЕПОВЕЦ']):
        return 'Вологда'
    if any(k in combined for k in ['ПЕНЗА']):
        return 'Пенза'
    if any(k in combined for k in ['КУРГАН']):
        return 'Курган'

    # Fallback regex
    m_city = re.search(r'(?:г\.|г\s+|город\s+)([А-Яа-яЁё\-]+)', t, re.IGNORECASE)
    if m_city:
        c = m_city.group(1).strip()
        return c.capitalize() if c.isupper() or c.islower() else c
    if t:
        parts = [x.strip() for x in t.split(',') if x.strip()]
        if parts:
            first = re.sub(r'^(г\.|г\s+|город\s+)', '', parts[0], flags=re.IGNORECASE).strip()
            if len(first.split()) <= 2 and not any(ch.isdigit() for ch in first):
                return first.capitalize() if first.isupper() or first.islower() else first
    return 'Другие города'

def main():
    json_path = os.path.join(ROOT, 'site', 'data.json')
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    sep_records = [x for x in data.get('sys_db_partners', []) if x.get('Month') == '2026-09']
    deals = [x for x in sep_records if x.get('Type') == 'Сделка']
    leads = [x for x in sep_records if x.get('Type') == 'Лид']
    print(f"Loaded: {len(deals)} deals, {len(leads)} leads for September 2026.")

    # 1. Read deal cities from Bitrix exports
    deals_files = [
        os.path.join(ROOT, 'raw_data', 'DEAL_20261002_971ed50c_6ac00c559314a.xls'),
        os.path.join(ROOT, 'raw_data', 'DEAL_20261007_e7742183_6ac5ece3ed91b.xls')
    ]
    deal_raw_cities = {}
    for df in deals_files:
        if not os.path.exists(df):
            continue
        m = read_tabular_file(df)
        for sname, rows in m:
            for r in rows:
                did = str(r.get('ID') or r.get('id') or '')
                city = str(r.get('Город. B2C') or r.get('ГОРОДB2C') or '').strip()
                if did and city and did not in deal_raw_cities:
                    deal_raw_cities[did] = city

    # 2. Read lead addresses
    lead_file = os.path.join(ROOT, 'raw_data', 'data (48).xlsx')
    lead_raw_addrs = {}
    if os.path.exists(lead_file):
        m_leads = read_tabular_file(lead_file)
        lead_rows = m_leads[0][1]
        for r in lead_rows:
            lid = str(r.get('Сумма id') or r.get('СУММАID') or '')
            cid = str(r.get('client_id') or r.get('CLIENT_ID') or '')
            addr = str(r.get('ADDRESS') or r.get('address') or '').strip()
            if lid and addr and lid not in lead_raw_addrs:
                lead_raw_addrs[lid] = addr
            if cid and addr and cid not in lead_raw_addrs:
                lead_raw_addrs[cid] = addr

    # 3. Default partner cities from registry
    registry = data.get('partners_registry', [])
    partner_default_city = {}
    for p in registry:
        cname = p.get('canonical_name')
        pid = p.get('partner_id')
        for o in p.get('oem_data', []):
            c = o.get('city')
            if c:
                clean_c = resolve_clean_city(c, cname)
                if clean_c != 'Другие города':
                    partner_default_city[cname] = clean_c
                    partner_default_city[pid] = clean_c
                    break

    # Aggregate by (City, Dealer, Brand)
    # data_matrix[(city, dealer, brand)] = {'leads': X, 'deals': Y}
    data_matrix = defaultdict(lambda: {'leads': 0, 'deals': 0})
    city_totals = defaultdict(lambda: {'leads': 0, 'deals': 0})
    dealer_totals = defaultdict(lambda: {'leads': 0, 'deals': 0})

    # Process DEALS
    for d in deals:
        dealer = d.get('Partner') or d.get('RawPartner') or 'Не указан'
        brand = d.get('Brand') or 'Не указан'
        did = str(d.get('DealId') or '')
        raw_city = deal_raw_cities.get(did, '')
        city = resolve_clean_city(raw_city, dealer)
        if city == 'Другие города' and dealer in partner_default_city:
            city = partner_default_city[dealer]
        data_matrix[(city, dealer, brand)]['deals'] += 1
        city_totals[city]['deals'] += 1
        dealer_totals[(city, dealer)]['deals'] += 1

    # Process LEADS
    for l in leads:
        dealer = l.get('Partner') or l.get('RawPartner') or 'Не указан'
        brand = l.get('Brand') or 'Не указан'
        lid = str(l.get('LeadId') or '')
        cid = str(l.get('ClientId') or '')
        raw_addr = lead_raw_addrs.get(lid) or lead_raw_addrs.get(cid) or ''
        city = resolve_clean_city(raw_addr, dealer)
        if city == 'Другие города' and dealer in partner_default_city:
            city = partner_default_city[dealer]
        data_matrix[(city, dealer, brand)]['leads'] += 1
        city_totals[city]['leads'] += 1
        dealer_totals[(city, dealer)]['leads'] += 1

    print(f"Matrix built: {len(data_matrix)} unique (City, Dealer, Brand) rows across {len(city_totals)} cities.")

    # Sort Cities: by deals desc, then leads desc, then name asc
    sorted_cities = sorted(city_totals.keys(), key=lambda c: (-city_totals[c]['deals'], -city_totals[c]['leads'], c))

    # Target path
    target_path = r'C:\Users\pc\Downloads\September_Deals_2026.xlsx'
    
    # Load existing or create new workbook
    if os.path.exists(target_path):
        wb = openpyxl.load_workbook(target_path)
    else:
        wb = openpyxl.Workbook()

    # Create new Sheet at index 0
    sheet_title = 'Города, дилеры и лиды'
    if sheet_title in wb.sheetnames:
        del wb[sheet_title]
    
    ws = wb.create_sheet(title=sheet_title, index=0)
    ws.views.sheetView[0].showGridLines = True

    # Typography & Styles
    font_family = 'Segoe UI'
    title_font = Font(name=font_family, size=15, bold=True, color='1F497D')
    sub_font = Font(name=font_family, size=10, italic=True, color='595959')
    
    header_font = Font(name=font_family, size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    
    city_header_font = Font(name=font_family, size=11, bold=True, color='1F497D')
    city_header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    
    subtotal_font = Font(name=font_family, size=10, bold=True, color='203764')
    subtotal_fill = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
    
    grand_total_font = Font(name=font_family, size=11, bold=True, color='000000')
    grand_total_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')

    data_font = Font(name=font_family, size=10)
    zebra_fill = PatternFill(start_color='FAFBFC', end_color='FAFBFC', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='D9D9D9')
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    thick_top = Side(border_style='thin', color='1F4E78')
    double_bottom = Side(border_style='double', color='1F4E78')
    total_border = Border(top=thick_top, bottom=double_bottom, left=thin_border_side, right=thin_border_side)

    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Title & Metadata
    ws['A1'] = 'Отчёт по лидам и сделкам: Города ➔ Дилеры ➔ Бренды авто'
    ws['A1'].font = title_font
    
    total_deals_cnt = sum(c['deals'] for c in city_totals.values())
    total_leads_cnt = sum(c['leads'] for c in city_totals.values())
    ws['A2'] = f'Период: Сентябрь 2026 г. | Передано лидов: {total_leads_cnt:,} | Сделок: {total_deals_cnt:,} | Городов: {len(sorted_cities)} | Строк воронки: {len(data_matrix)}'.replace(',', ' ')
    ws['A2'].font = sub_font

    headers = [
        '№ п/п',
        'Город',
        'Дилер / Холдинг',
        'Бренд автомобиля',
        'Передано лидов',
        'Сделки',
        'Конверсия (CR, %)'
    ]
    start_row = 4

    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_header
        c.border = data_border
    ws.row_dimensions[start_row].height = 28

    current_row = start_row + 1
    item_idx = 1
    subtotal_rows = []

    for city in sorted_cities:
        # Find all dealers in this city
        city_dealers = set(d for (c, d, b) in data_matrix.keys() if c == city)
        sorted_dealers = sorted(city_dealers, key=lambda d: (-dealer_totals[(city, d)]['deals'], -dealer_totals[(city, d)]['leads'], d))
        
        city_start_row = current_row

        for dealer in sorted_dealers:
            # Find all brands for this (city, dealer)
            dealer_brands = [b for (c, d, b) in data_matrix.keys() if c == city and d == dealer]
            sorted_brands = sorted(dealer_brands, key=lambda b: (-data_matrix[(city, dealer, b)]['deals'], -data_matrix[(city, dealer, b)]['leads'], b))

            for brand in sorted_brands:
                metrics = data_matrix[(city, dealer, brand)]
                leads_val = metrics['leads']
                deals_val = metrics['deals']

                ws.cell(row=current_row, column=1, value=item_idx).alignment = align_center
                ws.cell(row=current_row, column=2, value=city).alignment = align_left
                ws.cell(row=current_row, column=3, value=dealer).alignment = align_left
                ws.cell(row=current_row, column=4, value=brand).alignment = align_left
                
                # Leads
                c_leads = ws.cell(row=current_row, column=5, value=leads_val)
                c_leads.alignment = align_right
                c_leads.number_format = '#,##0'

                # Deals
                c_deals = ws.cell(row=current_row, column=6, value=deals_val)
                c_deals.alignment = align_right
                c_deals.number_format = '#,##0'

                # CR formula
                c_cr = ws.cell(row=current_row, column=7, value=f'=IF(E{current_row}>0, F{current_row}/E{current_row}, 0)')
                c_cr.alignment = align_right
                c_cr.number_format = '0.0%'

                is_even = (item_idx % 2 == 0)
                for col_idx in range(1, 8):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.font = data_font
                    cell.border = data_border
                    if is_even:
                        cell.fill = zebra_fill

                ws.row_dimensions[current_row].height = 20
                current_row += 1
                item_idx += 1

        city_end_row = current_row - 1

        # City Subtotal Row
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        sub_lbl = ws.cell(row=current_row, column=1, value=f'Итого по городу {city}')
        sub_lbl.alignment = Alignment(horizontal='right', vertical='center')

        sub_leads = ws.cell(row=current_row, column=5, value=f'=SUM(E{city_start_row}:E{city_end_row})')
        sub_leads.alignment = align_right
        sub_leads.number_format = '#,##0'

        sub_deals = ws.cell(row=current_row, column=6, value=f'=SUM(F{city_start_row}:F{city_end_row})')
        sub_deals.alignment = align_right
        sub_deals.number_format = '#,##0'

        sub_cr = ws.cell(row=current_row, column=7, value=f'=IF(E{current_row}>0, F{current_row}/E{current_row}, 0)')
        sub_cr.alignment = align_right
        sub_cr.number_format = '0.0%'

        for col_idx in range(1, 8):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = subtotal_font
            cell.fill = subtotal_fill
            cell.border = Border(top=Side(style='thin', color='B0C4DE'), bottom=Side(style='thin', color='B0C4DE'), left=thin_border_side, right=thin_border_side)

        ws.row_dimensions[current_row].height = 22
        subtotal_rows.append(current_row)
        current_row += 1

    # Grand Total Row
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    ws.cell(row=current_row, column=1, value='ИТОГО ПО ВСЕМ ГОРОДАМ').alignment = Alignment(horizontal='right', vertical='center')

    leads_sum_formula = '+'.join(f'E{r}' for r in subtotal_rows)
    deals_sum_formula = '+'.join(f'F{r}' for r in subtotal_rows)

    tot_leads = ws.cell(row=current_row, column=5, value=f'={leads_sum_formula}')
    tot_leads.alignment = align_right
    tot_leads.number_format = '#,##0'

    tot_deals = ws.cell(row=current_row, column=6, value=f'={deals_sum_formula}')
    tot_deals.alignment = align_right
    tot_deals.number_format = '#,##0'

    tot_cr = ws.cell(row=current_row, column=7, value=f'=IF(E{current_row}>0, F{current_row}/E{current_row}, 0)')
    tot_cr.alignment = align_right
    tot_cr.number_format = '0.0%'

    for col_idx in range(1, 8):
        c = ws.cell(row=current_row, column=col_idx)
        c.font = grand_total_font
        c.fill = grand_total_fill
        c.border = total_border
    ws.row_dimensions[current_row].height = 26

    # AutoFilter & Freeze Panes
    ws.auto_filter.ref = f'A{start_row}:G{current_row-1}'
    ws.freeze_panes = 'C5'

    # Column Widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 44
    ws.column_dimensions['D'].width = 24
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 16
    ws.column_dimensions['G'].width = 20

    # Save to Downloads and all target locations
    save_locations = [
        target_path,
        os.path.join(ROOT, 'September_Deals_2026.xlsx'),
        r'C:\Users\pc\OneDrive\Desktop\September_Deals_2026.xlsx'
    ]

    for sp in save_locations:
        try:
            wb.save(sp)
            print(f"Successfully saved to: {sp}")
        except Exception as e:
            print(f"Failed to save to {sp}: {e}")

if __name__ == '__main__':
    main()
