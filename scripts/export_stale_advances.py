import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_FILE = os.path.join(PROJECT_ROOT, 'data.json')
OUTPUT_ROOT = os.path.join(PROJECT_ROOT, 'Зависшие_авансы_от_20_дней.xlsx')
OUTPUT_SITE = os.path.join(PROJECT_ROOT, 'site', 'Зависшие_авансы_от_20_дней.xlsx')

def export_stale_advances():
    if not os.path.exists(DATA_FILE):
        print(f"Error: {DATA_FILE} not found")
        return

    with open(DATA_FILE, 'r', encoding='utf-8') as f:
        data = json.load(f)

    debtors = data.get('debtors', [])
    stale_20 = [d for d in debtors if d.get('aging_days', 0) >= 20]
    stale_20.sort(key=lambda x: (-x.get('aging_days', 0), x.get('company', '')))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Авансы от 20 дней"

    # Title Banner
    ws.merge_cells('A1:O1')
    title_cell = ws['A1']
    title_cell.value = f"РЕЕСТР ЗАВИСШИХ АВАНСОВ И ПРЕДОПЛАТ (ОТ 20 ДНЕЙ) — НА КОНТРОЛЬ КООРДИНАТОРАМ И КАМ"
    title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 40

    ws.merge_cells('A2:O2')
    sub_cell = ws['A2']
    sub_cell.value = f"Всего на контроле: {len(stale_20)} сделок с внесенным авансом без реализации ДКП | Для проверки дилерских центров и актуализации статусов в CRM"
    sub_cell.font = Font(name='Calibri', size=10, italic=True, color='CBD5E1')
    sub_cell.fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    sub_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 24

    headers = [
        "№",
        "ID Сделки",
        "ID Клиента",
        "Партнер / ДЦ",
        "Юрлицо в CRM",
        "Ответственный КАМ",
        "Дата аванса",
        "Дней зависания",
        "Марка",
        "Модель",
        "ВИН (VIN)",
        "Канал",
        "Менеджер сделки",
        "Текущая стадия",
        "Ссылка на CRM BackOffice"
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 28

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    
    # Cohort alert fills
    critical_fill = PatternFill(start_color='FFE4E6', end_color='FFE4E6', fill_type='solid') # >60 days
    warning_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')  # 31-60 days
    fresh_fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')    # 20-30 days

    for idx, d in enumerate(stale_20, 1):
        row_idx = header_row + idx
        ws.row_dimensions[row_idx].height = 22
        fill = zebra_fill if idx % 2 == 0 else white_fill
        days = d.get('aging_days', 0)

        did = str(d.get('deal_id') or '—')
        cid = str(d.get('client_id') or '—')
        backoffice_url = f"https://back.sberauto.com/crm/leads/{did}" if did != '—' else (f"https://back.sberauto.com/crm/clients/{cid}" if cid != '—' else "—")

        row_values = [
            idx,
            did,
            cid,
            d.get('company') or '—',
            d.get('raw_company') or '—',
            d.get('kam') or 'Не назначен',
            d.get('prepay_date') or '—',
            days,
            d.get('brand') or '—',
            d.get('model') or d.get('brand') or '—',
            d.get('vin') or '—',
            d.get('b2c') or '—',
            d.get('manager') or '—',
            d.get('stage') or 'В ожидании ДКП',
            backoffice_url
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            cell.fill = fill

            # Alignments
            if col_idx in [1, 2, 3, 7, 8, 12]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

            # Highlight aging column
            if col_idx == 8:
                if days > 60:
                    cell.fill = critical_fill
                    cell.font = Font(name='Calibri', size=10, bold=True, color='9F1239')
                elif days > 30:
                    cell.fill = warning_fill
                    cell.font = Font(name='Calibri', size=10, bold=True, color='92400E')
                else:
                    cell.fill = fresh_fill
                    cell.font = Font(name='Calibri', size=10, bold=True, color='065F46')

            # Link styling
            if col_idx == 15 and val.startswith('http'):
                cell.hyperlink = val
                cell.font = Font(name='Calibri', size=10, color='1D4ED8', underline='single')

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            v = str(cell.value or '')
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 22
    ws.column_dimensions['K'].width = 24
    ws.column_dimensions['O'].width = 40

    wb.save(OUTPUT_ROOT)
    wb.save(OUTPUT_SITE)
    print(f"[+] Успешно сформирован файл: {OUTPUT_ROOT} ({len(stale_20)} записей)")
    print(f"[+] Копия сохранена: {OUTPUT_SITE}")

if __name__ == '__main__':
    export_stale_advances()
